from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


OUTPUT = "/Users/rhee/AirConnect2/outputs/airconnect_backend_report/AirConnect_backend_report.docx"
DIAGRAM_DIR = Path("/Users/rhee/AirConnect2/outputs/airconnect_backend_report/diagrams")


def set_font(run, name="Calibri", east_asia="Malgun Gothic", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, name="Calibri", east_asia="Malgun Gothic", size=11, color=None, bold=None):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    style.font.size = Pt(size)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold


def set_paragraph_spacing(style, before=0, after=6, line=1.10):
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    set_style_font(styles["Normal"], size=10.8)
    set_paragraph_spacing(styles["Normal"], after=5, line=1.07)

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        set_style_font(styles[name], size=size, color=color, bold=True)
        set_paragraph_spacing(styles[name], before=before, after=after, line=1.10)


def add_text(paragraph, text, size=11, bold=False, color=None):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return run


def add_paragraph(doc, text="", style=None, align=None, keep_with_next=False):
    p = doc.add_paragraph(style=style)
    if text:
        add_text(p, text)
    if align is not None:
        p.alignment = align
    p.paragraph_format.keep_with_next = keep_with_next
    return p


def add_heading(doc, number, title, level=1):
    text = f"{number}. {title}" if number else title
    return add_paragraph(doc, text, style=f"Heading {level}", keep_with_next=True)


def set_cell_text(cell, text, bold=False, color=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_font(r, size=10, bold=bold, color=color)


def set_table_borders(table, color="B7C2CC"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def shade_cell(cell, color="F2F4F7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total = sum(widths)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_mar = tc_pr.find(qn("w:tcMar"))
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for side, value in [("top", "80"), ("bottom", "80"), ("start", "120"), ("end", "120")]:
                node = tc_mar.find(qn(f"w:{side}"))
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    tc_mar.append(node)
                node.set(qn("w:w"), value)
                node.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    for i, header in enumerate(headers):
        shade_cell(table.rows[0].cells[i])
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="1F4D78")
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            set_cell_text(row.cells[i], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def load_diagram_font(size, bold=False):
    candidates = [
        "/System/Library/Fonts/AppleSDGothicNeo.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    for path in candidates:
        try:
            return ImageFont.truetype(path, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def text_size(draw, text, font):
    bbox = draw.textbbox((0, 0), text, font=font)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def wrap_text(draw, text, font, max_width):
    words = text.split()
    lines = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if text_size(draw, candidate, font)[0] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_centered_text(draw, box, text, font, fill="1D2939", max_lines=3):
    x1, y1, x2, y2 = box
    lines = wrap_text(draw, text, font, (x2 - x1) - 34)
    if len(lines) > max_lines:
        lines = lines[: max_lines - 1] + [" ".join(lines[max_lines - 1 :])]
    line_height = text_size(draw, "가", font)[1] + 10
    total_height = line_height * len(lines)
    y = y1 + ((y2 - y1) - total_height) / 2
    for line in lines:
        width, height = text_size(draw, line, font)
        draw.text((x1 + ((x2 - x1) - width) / 2, y), line, font=font, fill=f"#{fill}")
        y += line_height


def draw_arrow(draw, start, end, color):
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=color, width=5)
    draw.polygon([(x2, y2), (x2 - 18, y2 - 11), (x2 - 18, y2 + 11)], fill=color)


def create_flow_diagram(filename, title, steps, accent):
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    path = DIAGRAM_DIR / filename
    width, height = 1800, 430
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font = load_diagram_font(36, bold=True)
    box_font = load_diagram_font(29)
    small_font = load_diagram_font(22)

    accent_color = f"#{accent}"
    muted_color = "#D9E2EC"
    box_fill = "#F8FBFE"
    title_width, _ = text_size(draw, title, title_font)
    draw.text(((width - title_width) / 2, 24), title, font=title_font, fill=accent_color)

    count = len(steps)
    margin_x = 68
    gap = 30
    box_width = (width - (margin_x * 2) - (gap * (count - 1))) / count
    box_height = 178
    y1 = 132
    y2 = y1 + box_height

    for index, step in enumerate(steps):
        x1 = margin_x + index * (box_width + gap)
        x2 = x1 + box_width
        box = (x1, y1, x2, y2)
        shadow = (x1 + 5, y1 + 7, x2 + 5, y2 + 7)
        draw.rounded_rectangle(shadow, radius=22, fill="#EEF3F8")
        draw.rounded_rectangle(box, radius=22, fill=box_fill, outline=muted_color, width=3)
        marker = f"{index + 1}"
        draw.ellipse((x1 + 18, y1 + 18, x1 + 58, y1 + 58), fill=accent_color)
        marker_width, marker_height = text_size(draw, marker, small_font)
        draw.text((x1 + 38 - marker_width / 2, y1 + 38 - marker_height / 2 - 1), marker, font=small_font, fill="#FFFFFF")
        draw_centered_text(draw, (x1 + 10, y1 + 48, x2 - 10, y2 - 10), step, box_font)
        if index < count - 1:
            draw_arrow(draw, (x2 + 8, (y1 + y2) / 2), (x2 + gap - 10, (y1 + y2) / 2), accent_color)

    draw.rounded_rectangle((margin_x, 346, width - margin_x, 388), radius=18, fill="#F2F4F7")
    note = "각 단계는 상태 검증, 중복 방지, 기록 보존을 전제로 순차 처리된다."
    note_width, note_height = text_size(draw, note, small_font)
    draw.text(((width - note_width) / 2, 356), note, font=small_font, fill="#475467")
    image.save(path, quality=95)
    return path


def ensure_flow_diagrams():
    specs = [
        (
            "flow_one_to_one_chat.png",
            "1대1 채팅 처리 흐름도",
            [
                "매칭 수락 또는 대화 진입",
                "참여 권한과 차단 관계 확인",
                "개인 대화 공간 생성 또는 복구",
                "메시지 저장과 최신 상태 갱신",
                "실시간 전달과 목록 갱신",
                "읽음 상태 반영과 알림 판단",
            ],
            "2E74B5",
            "그림 6. 1대1 채팅 처리 흐름도",
        ),
        (
            "flow_reward_ad.png",
            "광고 보상 처리 흐름도",
            [
                "광고 시청 요청과 보상 세션 생성",
                "광고 플랫폼의 완료 통지 수신",
                "전자 서명과 전송 정보 검증",
                "세션 상태와 중복 보상 확인",
                "티켓 지급 이력 기록",
                "잔액 반영과 결과 알림",
            ],
            "00856F",
            "그림 7. 광고 보상 처리 흐름도",
        ),
        (
            "flow_payment.png",
            "결제 처리 흐름도",
            [
                "결제 완료 정보 수신",
                "스토어 서버 검증",
                "주문 상태 잠금",
                "계정 식별과 중복 지급 확인",
                "상품 정책에 따른 지급 결정",
                "티켓 이력 기록과 잔액 반영",
            ],
            "9A6700",
            "그림 8. 결제 처리 흐름도",
        ),
        (
            "flow_ticket_ledger.png",
            "티켓 사용 이력 및 운영 점검 흐름도",
            [
                "지급 또는 차감 요청 발생",
                "사용자 잔액 기준 잠금",
                "정책 검증과 수량 계산",
                "티켓 사용 이력 저장",
                "잔액 일관성 반영",
                "통계와 운영 점검 자료 생성",
            ],
            "6A4BBC",
            "그림 9. 티켓 사용 이력 및 운영 점검 흐름도",
        ),
    ]
    return [(caption, create_flow_diagram(filename, title, steps, accent)) for filename, title, steps, accent, caption in specs]


def dashed_line(draw, xy, fill, width=4, dash=14, gap=10):
    x1, y1, x2, y2 = xy
    if x1 == x2:
        length = abs(y2 - y1)
        direction = 1 if y2 >= y1 else -1
        current = 0
        while current < length:
            start = y1 + direction * current
            end = y1 + direction * min(current + dash, length)
            draw.line((x1, start, x2, end), fill=fill, width=width)
            current += dash + gap
        return
    length = abs(x2 - x1)
    direction = 1 if x2 >= x1 else -1
    current = 0
    while current < length:
        start = x1 + direction * current
        end = x1 + direction * min(current + dash, length)
        draw.line((start, y1, end, y2), fill=fill, width=width)
        current += dash + gap


def draw_arrowhead(draw, x, y, direction, color):
    if direction >= 0:
        points = [(x, y), (x - 18, y - 10), (x - 18, y + 10)]
    else:
        points = [(x, y), (x + 18, y - 10), (x + 18, y + 10)]
    draw.polygon(points, fill=color)


def draw_sequence_participant(draw, x, y, label, font, accent):
    w, h = 236, 72
    rect = (x - w / 2, y, x + w / 2, y + h)
    draw.rounded_rectangle((rect[0] + 4, rect[1] + 5, rect[2] + 4, rect[3] + 5), radius=12, fill="#EFEAFB")
    draw.rounded_rectangle(rect, radius=12, fill="#F7F3FF", outline=accent, width=3)
    draw_centered_text(draw, rect, label, font, fill="111827", max_lines=2)


def draw_sequence_label(draw, x1, x2, y, text, font, fill="1D2939"):
    lines = wrap_text(draw, text, font, abs(x2 - x1) - 26)
    lines = lines[:2]
    line_h = text_size(draw, "가", font)[1] + 6
    top = y - 42 if len(lines) == 1 else y - 58
    for i, line in enumerate(lines):
        w, _ = text_size(draw, line, font)
        draw.text(((x1 + x2) / 2 - w / 2, top + i * line_h), line, font=font, fill=f"#{fill}")


def draw_sequence_arrow(draw, x1, x2, y, text, number, font, color="#333333", dashed=False):
    direction = 1 if x2 >= x1 else -1
    circle_x = x1
    draw.ellipse((circle_x - 18, y - 18, circle_x + 18, y + 18), fill="#2B2B2B")
    num_text = str(number)
    num_w, num_h = text_size(draw, num_text, font)
    draw.text((circle_x - num_w / 2, y - num_h / 2 - 1), num_text, font=font, fill="#FFFFFF")
    line_start = x1 + direction * 20
    line_end = x2 - direction * 22
    if dashed:
        dashed_line(draw, (line_start, y, line_end, y), fill=color, width=4)
    else:
        draw.line((line_start, y, line_end, y), fill=color, width=4)
    draw_arrowhead(draw, line_end + direction * 8, y, direction, color)
    draw_sequence_label(draw, line_start, line_end, y, text, font)


def draw_sequence_note(draw, x, y, text, font):
    w, h = 330, 82
    rect = (x - w / 2, y - h / 2, x + w / 2, y + h / 2)
    draw.rounded_rectangle(rect, radius=6, fill="#FFF6B7", outline="#C9A227", width=2)
    draw_centered_text(draw, rect, text, font, fill="3D2E00", max_lines=3)


def create_sequence_diagram(filename, title, lanes, steps, notes, groups, accent):
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    path = DIAGRAM_DIR / filename
    width = 2100
    row_h = 72
    top = 270
    bottom = 120
    height = top + len(steps) * row_h + bottom
    image = Image.new("RGB", (width, height), "#F1FCF8")
    draw = ImageDraw.Draw(image)
    title_font = load_diagram_font(34, bold=True)
    lane_font = load_diagram_font(28, bold=True)
    label_font = load_diagram_font(27)
    number_font = load_diagram_font(20, bold=True)
    accent_color = f"#{accent}"

    lane_left = 170
    lane_right = width - 170
    gap = (lane_right - lane_left) / (len(lanes) - 1)
    xs = [lane_left + i * gap for i in range(len(lanes))]

    title_w, _ = text_size(draw, title, title_font)
    draw.text(((width - title_w) / 2, 28), title, font=title_font, fill=accent_color)
    banner = (150, 86, width - 150, 134)
    draw.rounded_rectangle(banner, radius=4, fill="#FFF6B7", outline="#D5B543", width=2)
    banner_text = "주요 참여 주체별 요청, 검증, 저장, 반환 흐름"
    bw, bh = text_size(draw, banner_text, label_font)
    draw.text(((width - bw) / 2, 98), banner_text, font=label_font, fill="#4A3B00")

    for x, lane in zip(xs, lanes):
        draw_sequence_participant(draw, x, 142, lane, lane_font, accent_color)
        draw.line((x, 214, x, height - 78), fill="#9B7BEA", width=3)
        draw_sequence_participant(draw, x, height - 94, lane, lane_font, accent_color)

    for start_idx, end_idx, label in groups:
        y1 = top + (start_idx - 1) * row_h - 30
        y2 = top + (end_idx - 1) * row_h + 30
        x1, x2 = xs[0] - 36, xs[-1] + 36
        dashed_line(draw, (x1, y1, x2, y1), fill=accent_color, width=3)
        dashed_line(draw, (x2, y1, x2, y2), fill=accent_color, width=3)
        dashed_line(draw, (x2, y2, x1, y2), fill=accent_color, width=3)
        dashed_line(draw, (x1, y2, x1, y1), fill=accent_color, width=3)
        label_rect = (x1 - 4, y1 - 28, x1 + 104, y1 + 2)
        draw.rounded_rectangle(label_rect, radius=6, fill="#FFFFFF", outline=accent_color, width=2)
        lw, lh = text_size(draw, label, label_font)
        draw.text((x1 + 12, y1 - 25), label, font=label_font, fill=accent_color)

    for idx, (src, dst, text, dashed) in enumerate(steps, start=1):
        y = top + (idx - 1) * row_h
        draw_sequence_arrow(draw, xs[src], xs[dst], y, text, idx, number_font, dashed=dashed)

    for step_index, lane_index, note in notes:
        y = top + (step_index - 1) * row_h + 30
        draw_sequence_note(draw, xs[lane_index], y, note, label_font)

    image.save(path, quality=95)
    return path


def ensure_sequence_diagrams():
    specs = [
        (
            "sequence_one_to_one_chat.png",
            "1대1 채팅 시퀀스 흐름도",
            ["사용자", "모바일 앱", "요청 접수 계층", "채팅 처리 계층", "데이터 저장소", "알림 처리 계층"],
            [
                (0, 1, "매칭 수락 또는 대화 진입", False),
                (1, 2, "개인 대화 공간 생성 요청", False),
                (2, 3, "권한 검증과 대화 처리 호출", False),
                (3, 4, "참여자와 차단 관계 조회", False),
                (4, 3, "검증 데이터 반환", True),
                (3, 4, "대화 공간 생성 또는 복구", False),
                (3, 4, "메시지 저장과 최신 상태 갱신", False),
                (3, 1, "실시간 메시지와 목록 갱신 전달", False),
                (3, 5, "알림 필요 여부 판단", False),
                (5, 4, "알림 기록과 발송 대기 저장", False),
                (3, 2, "처리 결과 반환", True),
                (1, 0, "화면 갱신", True),
            ],
            [(4, 4, "참여 권한, 차단 여부, 방 상태 확인"), (9, 5, "수신자가 방을 보고 있으면 외부 발송 생략")],
            [(2, 7, "검증"), (8, 10, "알림")],
            "6F4CEB",
            "그림 7. 1대1 채팅 시퀀스 흐름도",
        ),
        (
            "sequence_reward_ad.png",
            "광고 보상 시퀀스 흐름도",
            ["사용자", "모바일 앱", "광고 플랫폼", "요청 접수 계층", "보상 처리 계층", "데이터 저장소"],
            [
                (0, 1, "광고 시청 진입", False),
                (1, 3, "보상 세션 발급 요청", False),
                (3, 4, "세션 생성 처리 호출", False),
                (4, 5, "보상 세션 저장", False),
                (5, 4, "세션 정보 반환", True),
                (1, 2, "광고 시청 완료", False),
                (2, 3, "완료 통지 전달", False),
                (3, 4, "전자 서명과 전송 정보 검증", False),
                (4, 5, "세션 상태와 중복 보상 조회", False),
                (4, 5, "티켓 지급 이력 저장", False),
                (4, 5, "잔액 반영", False),
                (4, 3, "보상 처리 결과 반환", True),
                (3, 1, "결과 동기화", True),
            ],
            [(8, 4, "서명 검증, 만료 여부, 중복 지급 확인"), (10, 5, "사용 이력 기준으로 잔액 변화 추적")],
            [(2, 5, "세션"), (7, 11, "보상")],
            "00856F",
            "그림 8. 광고 보상 시퀀스 흐름도",
        ),
        (
            "sequence_payment.png",
            "결제 시퀀스 흐름도",
            ["사용자", "모바일 앱", "스토어 검증 서버", "요청 접수 계층", "결제 처리 계층", "데이터 저장소"],
            [
                (0, 1, "상품 결제 완료", False),
                (1, 3, "결제 정보 전달", False),
                (3, 4, "결제 검증 처리 호출", False),
                (4, 2, "스토어 서버 검증 요청", False),
                (2, 4, "검증 결과 반환", True),
                (4, 5, "주문 상태 잠금", False),
                (4, 5, "계정 식별과 중복 지급 확인", False),
                (4, 5, "상품 정책에 따른 지급 결정", False),
                (4, 5, "티켓 지급 이력 저장", False),
                (4, 5, "잔액 반영", False),
                (4, 3, "결제 처리 결과 반환", True),
                (3, 1, "구매 결과와 잔액 갱신", True),
            ],
            [(4, 2, "구매 정보의 유효성과 환경 확인"), (7, 5, "동일 결제의 반복 지급 차단")],
            [(3, 8, "검증"), (9, 10, "지급")],
            "9A6700",
            "그림 9. 결제 시퀀스 흐름도",
        ),
    ]
    return [
        (caption, create_sequence_diagram(filename, title, lanes, steps, notes, groups, accent))
        for filename, title, lanes, steps, notes, groups, accent, caption in specs
    ]


def draw_entity_box(draw, x, y, w, h, title, items, fill, outline, title_font, item_font):
    draw.rounded_rectangle((x + 5, y + 7, x + w + 5, y + h + 7), radius=16, fill="#E8EEF5")
    draw.rounded_rectangle((x, y, x + w, y + h), radius=16, fill=fill, outline=outline, width=3)
    title_w, _ = text_size(draw, title, title_font)
    draw.text((x + (w - title_w) / 2, y + 14), title, font=title_font, fill="#111827")
    line_y = y + 52
    draw.line((x + 18, line_y, x + w - 18, line_y), fill=outline, width=2)
    current_y = y + 66
    for item in items:
        lines = wrap_text(draw, item, item_font, w - 32)
        for line in lines[:2]:
            draw.text((x + 18, current_y), line, font=item_font, fill="#344054")
            current_y += 27


def connect_boxes(draw, boxes, start, end, label, font, color="#667085"):
    x1, y1, w1, h1 = boxes[start]
    x2, y2, w2, h2 = boxes[end]
    sx, sy = x1 + w1 / 2, y1 + h1
    ex, ey = x2 + w2 / 2, y2
    if y2 < y1:
        sy = y1
        ey = y2 + h2
    draw.line((sx, sy, ex, ey), fill=color, width=3)
    draw.ellipse((ex - 6, ey - 6, ex + 6, ey + 6), fill=color)
    if label:
        tx, ty = (sx + ex) / 2, (sy + ey) / 2
        tw, th = text_size(draw, label, font)
        draw.rounded_rectangle((tx - tw / 2 - 8, ty - th / 2 - 5, tx + tw / 2 + 8, ty + th / 2 + 5), radius=8, fill="#FFFFFF")
        draw.text((tx - tw / 2, ty - th / 2 - 1), label, font=font, fill="#475467")


def create_project_erd():
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    path = DIAGRAM_DIR / "project_conceptual_erd.png"
    width, height = 2300, 1650
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font = load_diagram_font(42, bold=True)
    box_title_font = load_diagram_font(25, bold=True)
    item_font = load_diagram_font(21)
    rel_font = load_diagram_font(19)
    title = "프로젝트 스캔 기반 핵심 ERD"
    tw, _ = text_size(draw, title, title_font)
    draw.text(((width - tw) / 2, 30), title, font=title_font, fill="#1F4D78")
    subtitle = "실제 구현 모델을 기능 개념명으로 재구성한 데이터 관계 도식"
    sw, _ = text_size(draw, subtitle, item_font)
    draw.text(((width - sw) / 2, 86), subtitle, font=item_font, fill="#667085")

    palette = {
        "user": ("#EAF4FF", "#2E74B5"),
        "match": ("#F1F8FF", "#5B8DEF"),
        "chat": ("#F3F0FF", "#6F4CEB"),
        "group": ("#EFFFF8", "#00856F"),
        "notice": ("#FFF7E6", "#9A6700"),
        "ticket": ("#FFF4F4", "#C2410C"),
        "ops": ("#F5F5F5", "#667085"),
    }
    row_labels = [
        ("회원, 인증, 안전", 150, "#EAF4FF"),
        ("개인 매칭, 채팅", 380, "#F3F0FF"),
        ("그룹매칭", 610, "#EFFFF8"),
        ("알림, 푸시", 840, "#FFF7E6"),
        ("결제, 광고, 티켓", 1070, "#FFF4F4"),
        ("운영, 통계", 1300, "#F5F5F5"),
    ]
    for label, y, fill in row_labels:
        draw.rounded_rectangle((420, y - 42, 2220, y + 150), radius=18, fill=fill, outline="#E4E7EC", width=2)
        draw.text((438, y - 34), label, font=item_font, fill="#344054")

    boxes = {
        "사용자 정보": (70, 720, 300, 140, "user", ["계정 상태", "역할과 잔액 기준"]),
        "프로필 정보": (480, 150, 285, 120, "user", ["닉네임, 학교, 학과", "관심사와 취향"]),
        "인증, 학교 검증": (830, 150, 285, 120, "user", ["소셜 로그인", "학교 이메일 검증"]),
        "차단, 신고 기록": (1180, 150, 285, 120, "ops", ["차단 관계", "신고와 처리 상태"]),
        "추천 노출 기록": (480, 380, 285, 120, "match", ["후보 노출 이력", "중복 추천 방지"]),
        "개인 매칭 요청": (830, 380, 285, 120, "match", ["요청, 수락, 거절", "대화 연결 근거"]),
        "채팅방 정보": (1180, 380, 285, 120, "chat", ["개인, 그룹 대화 공간", "최신 메시지 기준"]),
        "참여자 정보": (1530, 380, 285, 120, "chat", ["방 참여 권한", "마지막 읽음 기준"]),
        "메시지 기록": (1880, 380, 285, 120, "chat", ["본문, 유형, 시각", "삭제와 읽음 상태"]),
        "팀방 정보": (480, 610, 285, 120, "group", ["모집 조건", "공개, 비공개 상태"]),
        "팀원, 준비 상태": (830, 610, 285, 120, "group", ["참여, 이탈, 추방", "준비 완료 여부"]),
        "그룹 매칭 결과": (1180, 610, 285, 120, "group", ["상대 팀 결정", "결과 상태"]),
        "최종 그룹방 정보": (1530, 610, 285, 120, "group", ["최종 참여자 묶음", "그룹 대화 연결"]),
        "알림 기록": (480, 840, 285, 120, "notice", ["수신자, 유형, 본문", "읽음, 삭제 상태"]),
        "장치, 발송 대기": (830, 840, 285, 120, "notice", ["푸시 장치", "재시도와 실패 상태"]),
        "푸시 이벤트 기록": (1180, 840, 285, 120, "notice", ["발송 결과", "열람과 실패 추적"]),
        "티켓 잔액": (480, 1070, 285, 120, "ticket", ["사용자별 보유 수량", "잠금 기준"]),
        "티켓 사용 이력": (830, 1070, 285, 120, "ticket", ["지급, 차감 기록", "중복 처리 방지"]),
        "결제 주문": (1180, 1070, 285, 120, "ticket", ["스토어 검증 결과", "중복 지급 방지"]),
        "광고 보상 기록": (1530, 1070, 285, 120, "ticket", ["보상 세션", "콜백 검증 결과"]),
        "활동, 통계 이벤트": (480, 1300, 285, 120, "ops", ["화면, 기능 이벤트", "운영 지표 근거"]),
        "공지, 점검 설정": (830, 1300, 285, 120, "ops", ["노출 기간, 우선순위", "버전과 점검 상태"]),
    }

    box_rects = {}
    for name, (x, y, w, h, kind, items) in boxes.items():
        fill, outline = palette[kind]
        draw_entity_box(draw, x, y, w, h, name, items, fill, outline, box_title_font, item_font)
        box_rects[name] = (x, y, w, h)

    bus_x = 410
    draw.line((box_rects["사용자 정보"][0] + box_rects["사용자 정보"][2], 790, bus_x, 790), fill="#667085", width=4)
    draw.line((bus_x, 210, bus_x, 1360), fill="#667085", width=4)

    def connect_user(target, label):
        x, y, w, h = box_rects[target]
        mid_y = y + h / 2
        draw.line((bus_x, mid_y, x, mid_y), fill="#667085", width=3)
        tw, th = text_size(draw, label, rel_font)
        draw.rounded_rectangle((bus_x + 20, mid_y - th - 16, bus_x + tw + 38, mid_y - 6), radius=8, fill="#FFFFFF")
        draw.text((bus_x + 28, mid_y - th - 13), label, font=rel_font, fill="#475467")

    for target, label in [
        ("프로필 정보", "1:1"),
        ("인증, 학교 검증", "1:N"),
        ("차단, 신고 기록", "1:N"),
        ("추천 노출 기록", "1:N"),
        ("팀방 정보", "1:N"),
        ("알림 기록", "1:N"),
        ("티켓 잔액", "1:1"),
        ("활동, 통계 이벤트", "1:N"),
    ]:
        connect_user(target, label)

    def connect_row(start, end, label):
        x1, y1, w1, h1 = box_rects[start]
        x2, y2, w2, h2 = box_rects[end]
        sx, sy = x1 + w1, y1 + h1 / 2
        ex, ey = x2, y2 + h2 / 2
        draw.line((sx, sy, ex, ey), fill="#667085", width=3)
        draw.ellipse((ex - 5, ey - 5, ex + 5, ey + 5), fill="#667085")
        tw, th = text_size(draw, label, rel_font)
        tx, ty = (sx + ex) / 2, (sy + ey) / 2
        draw.rounded_rectangle((tx - tw / 2 - 8, ty - th - 12, tx + tw / 2 + 8, ty - 2), radius=8, fill="#FFFFFF")
        draw.text((tx - tw / 2, ty - th - 9), label, font=rel_font, fill="#475467")

    def connect_bent(start, end, label):
        x1, y1, w1, h1 = box_rects[start]
        x2, y2, w2, h2 = box_rects[end]
        sx, sy = x1 + w1 / 2, y1 + h1
        ex, ey = x2 + w2 / 2, y2 + h2
        mid_y = max(sy, ey) + 34
        draw.line((sx, sy, sx, mid_y, ex, mid_y, ex, ey), fill="#667085", width=3)
        draw.ellipse((ex - 5, ey - 5, ex + 5, ey + 5), fill="#667085")
        tw, th = text_size(draw, label, rel_font)
        tx = (sx + ex) / 2
        draw.rounded_rectangle((tx - tw / 2 - 8, mid_y - th - 12, tx + tw / 2 + 8, mid_y - 2), radius=8, fill="#FFFFFF")
        draw.text((tx - tw / 2, mid_y - th - 9), label, font=rel_font, fill="#475467")

    for start, end, label in [
        ("프로필 정보", "인증, 학교 검증", "보완"),
        ("인증, 학교 검증", "차단, 신고 기록", "안전 기준"),
        ("추천 노출 기록", "개인 매칭 요청", "근거"),
        ("개인 매칭 요청", "채팅방 정보", "0:1"),
        ("채팅방 정보", "참여자 정보", "1:N"),
        ("참여자 정보", "메시지 기록", "1:N"),
        ("팀방 정보", "팀원, 준비 상태", "1:N"),
        ("팀원, 준비 상태", "그룹 매칭 결과", "조건 충족"),
        ("그룹 매칭 결과", "최종 그룹방 정보", "1:1"),
        ("알림 기록", "장치, 발송 대기", "1:N"),
        ("장치, 발송 대기", "푸시 이벤트 기록", "1:N"),
        ("티켓 잔액", "티켓 사용 이력", "1:N"),
        ("티켓 사용 이력", "결제 주문", "참조"),
        ("활동, 통계 이벤트", "공지, 점검 설정", "운영 판단"),
    ]:
        connect_row(start, end, label)
    connect_bent("티켓 사용 이력", "광고 보상 기록", "참조")

    legend = "관계 표기: 1:1 단일 관계, 1:N 다중 이력, 0:1 조건부 연결"
    lw, _ = text_size(draw, legend, item_font)
    draw.rounded_rectangle((width - lw - 120, 1530, width - 70, 1572), radius=12, fill="#F2F4F7")
    draw.text((width - lw - 96, 1538), legend, font=item_font, fill="#475467")
    image.save(path, quality=95)
    return path


def add_project_erd(doc):
    p = add_paragraph(doc, "그림 1. 프로젝트 스캔 기반 핵심 ERD", keep_with_next=True)
    set_font(p.runs[0], size=10, bold=True, color="1F4D78")
    picture_paragraph = doc.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(str(create_project_erd()), width=Inches(6.5))
    picture_paragraph.paragraph_format.space_after = Pt(8)


def add_flow_diagrams(doc):
    add_heading(doc, "", "주요 처리 흐름도", 2)
    add_paragraph(
        doc,
        "다음 도식은 사용자가 제공한 시퀀스 다이어그램 형식에 맞추어, 참여 주체별 요청과 검증, 저장, 반환 흐름을 세로 레인으로 재구성한 것이다. 제출 조건을 지키기 위해 실제 코드 식별자와 구체적인 요청 경로는 사용하지 않고, 백엔드 계층과 도메인 흐름을 개념명으로 표현하였다.",
    )
    for caption, path in ensure_sequence_diagrams():
        p = add_paragraph(doc, caption, keep_with_next=True)
        set_font(p.runs[0], size=10, bold=True, color="1F4D78")
        picture_paragraph = doc.add_paragraph()
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = picture_paragraph.add_run()
        run.add_picture(str(path), width=Inches(6.65))
        picture_paragraph.paragraph_format.space_after = Pt(8)


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = add_paragraph(doc, "AirConnect 백엔드 기능 설계 및 구현 보고서", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.runs[0], size=24, bold=True, color="0B2545")
    p.paragraph_format.space_after = Pt(18)

    p = add_paragraph(doc, "주제: 채팅, 알림, 그룹매칭, 통계 기능 중심", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.runs[0], size=14, bold=True, color="1F4D78")
    p.paragraph_format.space_after = Pt(36)

    for label in ["이름 : 정재민", "소속 : ______________________________", "학번 : ______________________________", "제출일 : ____________________________"]:
        p = add_paragraph(doc, label, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_font(p.runs[0], size=12)
        p.paragraph_format.space_after = Pt(10)

    for _ in range(5):
        doc.add_paragraph()
    p = add_paragraph(doc, "제출용 백엔드 설계 보고서", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.runs[0], size=11, color="555555")
    page_break(doc)


def add_toc(doc):
    add_heading(doc, "", "목차", 1)
    items = [
        "1. 프로젝트 개요",
        "2. 전체 백엔드 설계 개요",
        "3. 데이터 모델 및 ERD 분석",
        "4. 채팅 기능 설계 및 구현 설명",
        "5. 알림 기능 설계 및 구현 설명",
        "6. 그룹매칭 기능 설계 및 구현 설명",
        "7. 통계 기능 설계 및 구현 설명",
        "8. 기능 간 연계 흐름 및 주요 처리 흐름도",
        "9. 백엔드 설계상 주요 고려사항",
        "10. 구현 결과 및 기대 효과",
        "11. 결론",
    ]
    for item in items:
        p = add_paragraph(doc, item)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(5)
    page_break(doc)


SECTIONS = [
    (
        "1",
        "프로젝트 개요",
        [
            "AirConnect는 대학생이 같은 생활권과 관심사를 가진 사람을 탐색하고, 매칭 이후 대화로 자연스럽게 이어질 수 있도록 설계된 캠퍼스 기반 연결 서비스이다. 서비스의 핵심 문제의식은 단순한 프로필 노출이 아니라, 신뢰 가능한 사용자 정보와 상호 동의 기반의 연결 절차를 통해 관계 형성의 부담을 낮추는 데 있다. 사용자는 소셜 인증 또는 학교 기반 인증을 통해 서비스에 진입하고, 프로필과 취향 정보를 등록한 뒤 개인 매칭과 그룹매칭을 선택적으로 이용할 수 있다.",
            "제공된 프로젝트 설명 문서의 아키텍처 도식은 모바일 클라이언트, 서버 계층, 핵심 도메인, 저장소, 외부 서비스, 배포 환경이 서로 연결되는 구조를 보여준다. 이 도식의 의미는 백엔드가 단순히 요청을 받아 저장하는 역할에 머무르지 않고, 인증, 매칭, 채팅, 알림, 결제, 운영 관리가 하나의 연속된 서비스 흐름으로 결합되어 있다는 점이다. 즉 사용자의 한 번의 행동은 여러 도메인의 상태 변화와 외부 시스템 호출 및 운영 지표 축적으로 이어진다.",
            "사용자가 해결할 수 있는 문제는 크게 세 가지로 정리된다. 첫째, 사용자는 무작위성이 큰 만남 대신 학과, 성별, 관심사, 프로필 완성도와 같은 조건을 바탕으로 후보를 탐색할 수 있다. 둘째, 매칭 성사 후에는 별도의 외부 메신저로 이동하지 않고 서비스 내부에서 바로 대화를 시작할 수 있다. 셋째, 그룹 단위 만남에서는 팀 구성, 준비 확인, 대기열 진입, 최종 그룹 대화방 생성까지의 과정을 서버가 일관되게 관리함으로써 약속 형성의 복잡성을 줄인다.",
            "전체 백엔드 구조의 설계 방향은 계층형 아키텍처와 도메인 중심 책임 분리에 있다. 클라이언트와 직접 접하는 인터페이스 계층, 비즈니스 규칙을 처리하는 응용 계층, 상태를 표현하는 도메인 모델, 영속 데이터를 관리하는 저장소 계층이 분리되어 있다. 또한 인증과 보안, 예외 처리, 공통 응답, 실시간 통신 설정, 운영 로그와 같은 횡단 관심사는 별도의 공통 영역으로 다루어 핵심 기능이 자신의 책임에 집중할 수 있도록 설계되었다.",
            "본 보고서는 프로젝트 전체 기능 중 특히 채팅, 알림, 그룹매칭, 통계 기능을 중심으로 설명한다. 네 기능은 AirConnect 서비스 경험의 중심축이다. 그룹매칭은 사용자를 최종 관계 형성 단계로 이끄는 흐름을 담당하고, 채팅은 매칭 결과가 실제 소통으로 전환되는 지점을 담당한다. 알림은 사용자가 상태 변화를 놓치지 않도록 서비스 흐름을 연결하며, 통계는 사용자 활동과 운영 상태를 분석 가능한 정보로 전환한다. 이 네 기능은 독립된 기능이면서 동시에 서로 강하게 연결된 백엔드 구조의 사례이다.",
        ],
        None,
    ),
    (
        "2",
        "전체 백엔드 설계 개요",
        [
            "AirConnect 백엔드는 계층형 구조를 기반으로 한다. 클라이언트 요청은 인증 및 권한 검사를 거친 뒤 도메인별 처리 흐름으로 전달되고, 각 도메인은 필요한 데이터 검증과 상태 변경을 수행한다. 상태 변경이 필요한 기능은 트랜잭션 단위로 처리되어 데이터 정합성을 유지하고, 조회 중심 기능은 읽기 전용 흐름으로 분리되어 불필요한 변경을 방지한다. 이러한 구조는 기능 확장 시에도 기존 책임 경계를 유지할 수 있게 한다.",
            "도메인 중심 설계의 핵심은 서비스의 주요 명사를 데이터와 행위의 단위로 분리한 점이다. 사용자 정보, 프로필 정보, 매칭 요청 기록, 채팅방 정보, 메시지 기록, 알림 기록, 팀방 정보, 준비 상태, 최종 그룹방 정보, 티켓 사용 이력, 통계 이벤트가 각각 독립된 의미를 갖는다. 각 데이터는 단순 저장 대상이 아니라 서비스 상태 전이를 설명하는 근거로 작동한다.",
            "데이터베이스 기반 상태 관리는 AirConnect의 중요한 설계 선택이다. 매칭 요청의 대기, 수락, 거절, 그룹매칭의 모집, 준비 확인, 대기열 진입, 매칭 완료, 종료, 알림의 생성, 읽음, 삭제, 티켓의 지급, 차감은 모두 명확한 상태값과 이력으로 관리된다. 이를 통해 사용자가 같은 요청을 반복하거나 네트워크가 불안정한 상황에서도 서버가 최종 상태를 기준으로 일관된 응답을 제공할 수 있다.",
            "실시간성과 비동기성이 필요한 기능은 별도의 고려가 적용되었다. 채팅은 메시지 저장 이후 실시간 브로드캐스팅이 이어지는 구조이며, 알림은 원본 기록과 실제 푸시 발송을 분리하여 외부 서비스 장애에 대응한다. 그룹매칭은 대기열을 통해 여러 팀을 순서대로 처리하고, 백그라운드 작업이 대기열을 주기적으로 복구하거나 비워 안정성을 높인다. 통계는 클라이언트 이벤트와 서버 이벤트를 모두 수집하여 운영 판단에 활용한다.",
            "기능 간 연계 구조는 서비스의 완성도를 좌우한다. 개인 매칭이 수락되면 채팅방이 생성되거나 복구되고, 그 결과는 알림으로 전달된다. 그룹매칭이 성사되면 임시 팀방은 정리되고 최종 그룹 채팅방이 생성되며, 이 흐름은 실시간 이벤트와 푸시 알림, 통계 이벤트로 확장된다. 사용자의 메시지 전송은 채팅 기록과 채팅방 목록 갱신뿐 아니라 상대방 알림과 미읽음 수 계산으로 이어진다. 이러한 연계는 단일 기능 구현을 넘어 서비스 전체 흐름을 구성한다.",
        ],
        {
            "headers": ["설계 축", "적용 방식", "의미"],
            "rows": [
                ["책임 분리", "계층별 역할 분담", "기능 변경의 영향 범위 축소"],
                ["상태 관리", "관계형 데이터 기반 상태 전이", "반복 요청과 예외 상황에 대한 일관성 확보"],
                ["실시간 처리", "메시지 발행과 구독 구조", "대화와 매칭 상태 변화의 즉시 전달"],
                ["비동기 처리", "알림 발송 대기열과 재시도", "외부 서비스 장애에 대한 회복력 확보"],
                ["운영 분석", "활동 이벤트와 집계 지표", "서비스 운영 판단을 위한 근거 제공"],
            ],
            "widths": [1800, 3300, 4260],
        },
    ),
    (
        "3",
        "데이터 모델 및 ERD 분석",
        [
            "제공된 ERD 도식과 프로젝트 구현을 함께 검토하면, 전체 데이터 모델은 회원과 인증, 개인 매칭과 채팅, 그룹매칭, 알림과 푸시, 티켓과 결제, 광고 보상, 운영과 분석으로 나뉜다. 아래 핵심 ERD는 실제 프로젝트 스캔 결과를 바탕으로 주요 영속 모델을 코드 식별자가 아닌 기능 개념명으로 재구성한 것이다. 이는 AirConnect의 데이터 모델이 단순한 사용자 중심 구조가 아니라, 서비스 흐름별 상태와 이력을 추적하는 구조임을 보여준다.",
            "회원, 인증, 안전 영역에서는 사용자 정보가 프로필 정보, 학교 검증 정보, 인증 관련 정보, 신고 기록, 차단 기록과 연결된다. 이 관계는 인증과 사용자 상태가 이후 모든 기능의 선행 조건임을 의미한다. 예를 들어 채팅방 생성, 매칭 후보 추출, 그룹매칭 참여는 모두 사용자 상태와 차단 관계를 확인한 뒤 진행된다. 따라서 사용자 모델은 단순한 계정 정보가 아니라 서비스 이용 가능성을 판단하는 기준점이다.",
            "개인 매칭, 채팅 영역에서는 추천 노출 기록, 매칭 요청 기록, 채팅방 정보, 참여자 정보, 메시지 기록이 순차적으로 연결된다. 추천 노출 기록은 사용자가 실제로 본 후보에게만 요청을 보낼 수 있도록 하는 근거이며, 매칭 요청 기록은 수락 또는 거절 흐름을 관리한다. 매칭이 수락되면 채팅방 정보가 생성되거나 기존 관계와 연결되고, 참여자 정보와 메시지 기록이 대화의 영속성을 담당한다.",
            "그룹매칭 영역에서는 임시 팀방 정보, 팀원 정보, 준비 상태, 매칭 결과, 최종 그룹방 정보가 연결된다. 이 구조는 그룹매칭이 한 번의 무작위 추첨이 아니라 팀 형성, 준비 확인, 대기열 진입, 상대 팀 결정, 최종 대화방 생성이라는 단계적 상태 전이로 이루어짐을 보여준다. 각 단계의 데이터가 별도로 존재하기 때문에 서버는 사용자가 이탈하거나 네트워크 요청이 중복될 때도 현재 상태를 기준으로 흐름을 복구할 수 있다.",
            "알림, 푸시 영역은 알림 기록, 사용자별 알림 설정, 푸시 장치 정보, 발송 대기 기록, 푸시 이벤트 기록의 관계로 구성된다. 이 관계는 알림을 단순 메시지 표시로 보지 않고, 생성, 사용자 설정 반영, 발송 대상 결정, 외부 발송, 수신 또는 열람 이벤트 추적까지의 전체 흐름으로 설계했음을 의미한다.",
            "티켓, 결제, 광고 보상 영역은 사용자의 잔액 정보와 지급, 차감 기록이 결제 주문, 광고 보상 세션, 보상 검증 기록과 연결되는 구조를 가진다. 이는 AirConnect가 티켓을 단순 숫자 필드로만 관리하지 않고, 사용 이력을 통해 잔액 변화의 원인과 결과를 추적한다는 점에서 중요하다. 이러한 이력 관리 구조는 중복 지급 방지, 환불 처리, 운영 점검에 필요한 근거가 된다.",
            "운영, 분석 도식은 사용자 활동과 운영 데이터를 별도의 집계 대상으로 다룬다. 사용자 활동, 매칭 성사, 채팅 메시지, 알림 발송 상태, 신고 처리, 운영 공지, 데이터 정합성 점검 결과가 운영자의 판단 자료로 전환된다. 결과적으로 ERD는 AirConnect 백엔드가 기능 구현뿐 아니라 운영 가능한 서비스 구조를 지향하고 있음을 보여준다.",
        ],
        {
            "headers": ["개념 영역", "주요 관계", "설계 의미"],
            "rows": [
                ["회원, 인증", "사용자 정보와 프로필, 인증, 안전 기록", "모든 기능의 이용 가능성 판단"],
                ["채팅", "채팅방 정보와 참여자, 메시지 기록", "대화 영속성과 읽음 상태 관리"],
                ["그룹매칭", "팀방, 준비 상태, 매칭 결과, 최종 방", "단계별 상태 전이와 복구 가능성"],
                ["알림", "알림 기록, 설정, 장치, 발송 대기", "인앱 기록과 외부 발송의 분리"],
                ["통계", "활동 이벤트와 집계 데이터", "서비스 운영 판단의 근거 제공"],
            ],
            "widths": [1800, 3500, 4060],
        },
    ),
    (
        "4",
        "채팅 기능 설계 및 구현 설명",
        [
            "채팅 기능의 목적은 매칭 결과를 실제 소통으로 전환하는 것이다. 개인 매칭이 수락되거나 그룹매칭이 최종 성사되면 사용자는 별도 외부 채널로 이동하지 않고 서비스 내부에서 대화를 시작할 수 있다. 이를 위해 채팅 기능은 채팅방 생성, 참여자 관리, 메시지 저장, 메시지 조회, 읽음 처리, 실시간 전달, 목록 갱신, 알림 연동을 하나의 흐름으로 관리한다.",
            "대화가 생성되는 흐름은 채팅방 정보와 참여자 정보의 결합으로 시작된다. 개인 채팅은 두 사용자 간 관계를 기준으로 중복 생성을 방지하고, 이미 존재하는 대화 공간이 있으면 이를 재사용하거나 복구한다. 그룹 채팅은 여러 참여자를 한 번에 등록할 수 있도록 설계되어 임시 팀방과 최종 그룹방 모두에 사용된다. 참여자 정보에는 사용자가 언제 방에 들어왔는지, 어떤 메시지까지 읽었는지, 현재 노출 가능한 참여자인지가 저장된다.",
            "메시지 저장 구조는 실시간 전달보다 먼저 데이터 영속성을 확보하는 방향으로 설계되어 있다. 사용자가 메시지를 보내면 서버는 먼저 접근 권한과 차단 관계를 검증하고, 메시지 내용을 저장한 뒤 채팅방의 마지막 메시지 정보를 갱신한다. 이후 실시간 통신 채널로 메시지를 발행하고, 채팅방 목록 갱신 이벤트와 상대방 알림 생성 흐름을 수행한다. 이 순서는 실시간 발행이 실패하더라도 저장된 메시지를 조회를 통해 복구할 수 있게 한다.",
            "읽음 처리는 채팅 기능에서 중요한 부분이다. AirConnect는 참여자별 마지막 읽음 기준을 저장하고, 메시지별 미읽음 수를 계산한다. 개인 채팅에서는 상대방이 읽지 않은 경우 미읽음 수가 남고, 그룹 채팅에서는 아직 읽지 않은 참여자 수가 미읽음 수로 표현된다. 사용자가 채팅방을 조회하거나 실시간 구독 상태에 들어가거나 명시적으로 읽음 처리를 요청하는 경우 모두 같은 읽음 동기화 규칙을 거치도록 설계되어, 클라이언트 진입 경로가 달라도 결과가 일관된다.",
            "실시간 전달과 데이터 영속성의 균형은 저장 후 발행 구조에서 드러난다. 메시지는 관계형 데이터베이스에 저장되어 조회 가능성을 보장하고, 실시간 통신은 메시지 도착, 읽음 이벤트, 채팅방 목록 갱신을 빠르게 전달한다. 동시에 사용자가 이미 해당 채팅방을 보고 있으면 상대방 알림을 불필요하게 보내지 않는 방식으로 사용 경험과 푸시 발송 비용을 함께 고려한다.",
            "제공된 시퀀스 도식의 흐름을 자연어로 해석하면 다음과 같다. 사용자가 매칭 요청을 수락하면 서버는 두 사용자의 관계를 확인하고 개인 대화 공간을 만들거나 기존 공간을 복구한다. 이후 양쪽 사용자에게 대화 공간 접근 정보가 전달된다. 그룹매칭의 경우에는 팀 구성과 준비 확인이 끝난 뒤 두 팀이 매칭되고, 기존 임시 팀방은 정리되며 새로운 최종 그룹 대화 공간이 생성된다. 채팅 기능은 이 두 흐름의 종착점이자 이후 사용자 관계가 지속되는 장소이다.",
            "예외 처리 측면에서는 접근 권한이 없는 사용자가 메시지를 조회하거나 보낼 수 없도록 방 참여 여부를 검증한다. 차단 관계가 존재하면 대화 생성과 메시지 전송이 제한된다. 또한 메시지 삭제는 발신자 본인만 수행할 수 있고, 삭제된 메시지는 기록을 완전히 제거하기보다 표시 상태를 바꾸는 방식으로 처리되어 운영 관점의 추적 가능성을 유지한다.",
        ],
        None,
    ),
    (
        "5",
        "알림 기능 설계 및 구현 설명",
        [
            "알림 기능의 목적은 사용자가 서비스 상태 변화를 놓치지 않도록 하는 것이다. 매칭 요청 도착, 요청 수락 또는 거절, 그룹매칭 팀원 합류, 준비 상태 변경, 팀방 해산, 그룹매칭 성사, 채팅 메시지 도착, 운영 공지, 일정 리마인드와 같은 이벤트는 사용자가 즉시 인지해야 하는 서비스 변화이다. 알림 기능은 이러한 이벤트를 인앱 알림 기록과 외부 푸시 발송으로 나누어 처리한다.",
            "알림이 발생하는 조건은 도메인 이벤트와 연결되어 있다. 개인 매칭에서는 요청자가 상대에게 요청을 보낼 때와 상대가 수락 또는 거절할 때 알림이 생성된다. 채팅에서는 수신자가 현재 해당 방을 보고 있지 않을 때 메시지 알림이 생성된다. 그룹매칭에서는 팀원이 합류하거나 이탈할 때, 준비 상태가 바뀔 때, 모든 팀원이 준비되어 방장이 매칭을 시작할 수 있을 때, 최종 그룹매칭이 성사될 때 알림이 발생한다. 운영 기능에서는 공지와 시스템 안내가 알림으로 연결된다.",
            "알림 저장 구조는 원본 기록과 발송 대기 기록을 분리한다. 먼저 알림 원본이 생성되고, 사용자 설정에 따라 인앱 알림함에 노출될지와 외부 푸시로 발송될지가 결정된다. 사용자가 푸시를 허용하지 않았거나 방해금지 시간대에 있는 경우에도 원본 알림은 남을 수 있지만, 외부 푸시 발송은 제한된다. 이 구조는 사용자 경험을 존중하면서도 서비스 내부 기록을 유지하는 방식이다.",
            "발송 대기 구조는 비동기 처리의 핵심이다. 외부 푸시 발송은 네트워크 상태와 외부 서비스 응답에 영향을 받기 때문에, 원본 알림 저장과 같은 트랜잭션 안에서 즉시 성공해야 하는 작업으로 두지 않는다. 서버는 발송 가능한 항목을 일정 단위로 점유하고, 성공, 재시도, 실패, 정책상 제외와 같은 상태로 전이한다. 오래 처리 중인 항목은 복구되어 다시 시도될 수 있으므로, 일시 장애 이후에도 발송 흐름이 멈추지 않는다.",
            "사용자가 알림을 확인하는 흐름은 커서 기반 조회, 미읽음 수 조회, 단건 읽음 처리, 전체 읽음 처리, 삭제로 구성된다. 알림함은 채팅 메시지 알림처럼 별도 화면에서 직접 처리되는 항목을 제외하여 사용자가 운영, 매칭, 그룹매칭 관련 상태 변화를 더 명확히 볼 수 있게 한다. 삭제는 물리적 제거보다 사용자 화면에서 숨기는 방식으로 처리되어 데이터 추적성과 사용자 경험을 동시에 고려한다.",
            "중복 방지 역시 중요한 설계 요소이다. 같은 이벤트가 네트워크 재시도나 중복 호출로 여러 번 발생할 수 있으므로, 알림은 이벤트별 중복 방지 기준을 사용하여 이미 생성된 기록을 재사용할 수 있다. 채팅 메시지의 경우 짧은 시간 안에 같은 채팅방에서 반복되는 푸시를 묶는 정책을 적용하여 모바일 환경에서 과도한 알림을 줄인다.",
            "비동기 처리와 이벤트 기반 설계 관점에서 알림 기능은 다른 도메인의 상태 변화를 사용자에게 전달하는 연결 계층이다. 매칭과 그룹매칭이 상태를 바꾸면 알림 기능이 이를 사용자 인지로 전환하고, 채팅 기능은 새 메시지를 알림으로 확장한다. 운영자는 발송 성공률, 실패 사유, 오래 대기 중인 항목을 확인할 수 있으므로 알림 기능은 사용자 경험뿐 아니라 서비스 운영 안정성과도 직접 연결된다.",
        ],
        None,
    ),
    (
        "6",
        "그룹매칭 기능 설계 및 구현 설명",
        [
            "그룹매칭 기능의 목적은 여러 사용자가 팀을 구성한 뒤 상대 팀과 연결되어 최종 그룹 대화로 이동하는 흐름을 제공하는 것이다. 개인 매칭이 사용자 두 명의 상호 선택에 초점을 둔다면, 그룹매칭은 팀 단위의 준비 상태와 상대 팀 조건을 함께 고려해야 한다. 따라서 그룹매칭은 팀방 생성, 팀원 참여, 준비 확인, 대기열 진입, 상대 팀 매칭, 최종 그룹방 생성이라는 단계적 절차로 설계되었다.",
            "팀방 생성 단계에서는 방장이 팀 이름, 팀 크기, 팀 성별 조건, 상대 팀 조건, 공개 여부를 지정한다. 방이 생성되면 임시 대화 공간도 함께 만들어져 팀원 간 조율이 가능하다. 공개 팀방은 조건에 맞는 사용자가 목록에서 찾아 들어올 수 있고, 비공개 팀방은 초대 코드를 통해 참여할 수 있다. 이 구조는 공개 모집과 지인 기반 팀 구성이라는 두 가지 사용 방식을 모두 지원한다.",
            "매칭 조건과 사용자 정보의 활용 방식은 공정성과 정확성에 영향을 준다. 서버는 사용자의 프로필 성별과 팀 성별 조건이 맞는지 확인하고, 같은 사용자가 동시에 여러 활성 팀방에 참여하지 못하도록 제한한다. 팀 크기가 충족되면 준비 확인 단계로 전환되며, 모든 팀원이 준비를 완료해야 대기열 진입이 가능하다. 준비 시점에는 필요한 티켓을 보유했는지도 확인하여 매칭 완료 직전에 실패할 가능성을 줄인다.",
            "대기열 관리는 그룹매칭의 핵심이다. 방장이 매칭을 시작하면 팀방은 대기 상태로 전환되고, 팀 크기별 대기열에 등록된다. 서버는 대기열을 순서대로 읽으면서 오래되었거나 상태가 맞지 않는 항목을 정리하고, 조건이 맞는 두 팀을 찾는다. 이때 분산 잠금을 사용하여 여러 작업자가 동시에 같은 팀을 매칭하지 않도록 하고, 대기열 토큰과 복구 절차를 통해 서버 재시작 이후에도 대기 상태를 동기화할 수 있다.",
            "매칭 결과가 결정되면 두 팀의 상태는 매칭 완료 단계로 전환된다. 이후 두 팀의 모든 사용자가 포함된 최종 그룹 대화 공간이 생성되고, 임시 팀방의 참여 상태와 준비 상태는 정리된다. 티켓은 최종 그룹방 생성 이후 차감되며, 사용자별 잔액을 잠근 상태에서 처리하여 동시성 문제를 줄인다. 최종 결과는 실시간 이벤트와 알림으로 사용자에게 전달된다.",
            "예외 상황도 단계별로 관리된다. 팀원이 나가면 준비 상태가 초기화되고, 방장이 팀을 해산하면 대기열에서 제거된다. 팀이 가득 차지 않았거나 준비가 완료되지 않았거나 상대 팀 조건이 맞지 않으면 대기열 진입 또는 매칭이 진행되지 않는다. 추방된 사용자의 재입장, 방장의 자기 추방, 종료된 팀방 접근, 대기열 중복 등록과 같은 예외도 별도로 제한된다.",
            "서비스 관점에서 그룹매칭은 공정성, 정확성, 확장성을 모두 고려한다. 공정성은 대기열 순서와 조건 기반 탐색에서 확보되고, 정확성은 팀원 수, 준비 상태, 상태 전이를 매칭 직전까지 검증하는 방식으로 확보된다. 확장성은 팀 크기별 대기열, 백그라운드 처리, 대기열 복구, 실시간 이벤트 발행 구조를 통해 확보된다. 이러한 설계는 향후 학교, 학과, 관심사, 시간대, 선호 조건이 추가되더라도 기존 흐름 위에 조건을 확장할 수 있는 기반이 된다.",
        ],
        None,
    ),
    (
        "7",
        "통계 기능 설계 및 구현 설명",
        [
            "통계 기능의 목적은 사용자 활동과 서비스 상태를 운영자가 이해할 수 있는 지표로 전환하는 것이다. 매칭 서비스는 기능이 많아질수록 단순한 총 사용자 수만으로는 서비스 품질을 판단하기 어렵다. AirConnect의 통계 기능은 사용자 규모, 당일 활동, 성별 구성, 매칭 성사, 학과별 요청 흐름, 그룹매칭 퍼널, 알림 발송 상태, 신고 처리와 같은 데이터를 운영 관점에서 해석할 수 있도록 구성되어 있다.",
            "주요 집계 데이터는 서비스 흐름과 직접 연결되어 있다. 사용자 정보에서는 전체 활성 가입자와 당일 활동 사용자가 집계된다. 프로필 정보에서는 성별 비율이 계산된다. 개인 매칭 기록에서는 수락된 연결 수가, 그룹매칭 기록에서는 활성 또는 종료된 최종 그룹방 수가 매칭 성과로 집계된다. 매칭 요청 대상의 학과 정보는 인기 요청 학과 순위로 전환되어 사용자 관심이 집중되는 영역을 보여준다.",
            "통계 기능은 단순 누적 수치뿐 아니라 행동 흐름을 분석한다. 사용자가 추천을 새로고침하고, 매칭 요청을 보내고, 요청이 수락되며, 채팅방이 생성되는 과정을 퍼널로 볼 수 있다. 그룹매칭에서는 팀방 생성, 팀방 참여, 준비 완료, 대기열 진입, 매칭 성공, 최종 그룹 대화방 생성으로 이어지는 단계를 분석할 수 있다. 이러한 퍼널 구조는 특정 단계에서 사용자가 많이 이탈하는지 파악할 수 있게 한다.",
            "사용자 활동 데이터는 클라이언트에서 전송되는 화면 진입, 세션 시작과 종료, 푸시 열람 같은 이벤트와 서버 내부에서 기록하는 매칭, 그룹매칭, 티켓, 결제 이벤트가 결합되어 구성된다. 서버 이벤트는 주요 기능 흐름이 실제로 완료된 시점에 기록되므로, 클라이언트 로그만으로 확인하기 어려운 백엔드 처리 결과를 보완한다. 이 방식은 통계의 신뢰도를 높인다.",
            "데이터 집계 기준은 시간 범위와 상태 기준으로 나뉜다. 운영자는 최근 일정 기간 동안의 매칭 흐름, 알림 발송 성공률, 실패 사유, 신고 처리 현황, 데이터 정합성 점검 결과를 확인할 수 있다. 통계는 실시간성보다는 운영 판단에 필요한 일관성과 해석 가능성을 우선한다. 따라서 읽기 전용 트랜잭션과 집계 쿼리를 중심으로 구성되어 기능 처리 흐름에 불필요한 부담을 주지 않는다.",
            "관리자 또는 서비스 운영 관점에서 통계 기능의 의미는 크다. 매칭 성사 수는 서비스의 핵심 가치가 실제로 발생했는지 보여주고, 활동 사용자 수는 서비스의 현재 활력을 보여준다. 학과별 요청 순위는 사용자 선호와 서비스 홍보 전략을 결정하는 근거가 된다. 알림 발송 상태와 실패 원인은 모바일 사용자 경험의 장애를 조기에 발견하는 지표가 된다. 신고와 제재 관련 지표는 안전한 커뮤니티 운영의 기준을 제공한다.",
        ],
        None,
    ),
    (
        "8",
        "기능 간 연계 흐름 및 주요 처리 흐름도",
        [
            "채팅과 알림의 연계는 사용자가 대화 흐름을 놓치지 않도록 하는 데 목적이 있다. 사용자가 메시지를 보내면 서버는 메시지를 저장하고 실시간으로 채팅방 참여자에게 전달한다. 수신자가 현재 해당 채팅방을 보고 있으면 읽음 상태가 즉시 반영되고 외부 푸시는 생략될 수 있다. 반대로 수신자가 방을 보고 있지 않으면 알림 기록이 생성되고, 사용자 설정에 따라 외부 푸시 발송 대기열에 등록된다. 이 흐름은 데이터 저장, 실시간 전달, 알림 발행이 순차적으로 결합된 구조이다.",
            "그룹매칭과 알림의 연계는 상태 전이를 사용자에게 알려주는 방식으로 작동한다. 팀원이 합류하면 기존 팀원에게 합류 알림이 전달되고, 팀이 가득 차면 준비 확인 알림이 생성된다. 팀원별 준비 상태가 바뀌면 다른 팀원에게 알림이 전달되며, 모든 팀원이 준비되면 방장이 매칭을 시작할 수 있음을 알 수 있다. 최종 매칭이 성사되면 모든 참여자에게 최종 그룹 대화 공간으로 이동하라는 알림이 생성된다.",
            "그룹매칭 결과와 통계 데이터의 연계는 운영 분석을 가능하게 한다. 팀방이 생성될 때, 팀원이 참여할 때, 준비 상태가 바뀔 때, 대기열에 진입할 때, 최종 매칭이 완료될 때 서버는 활동 이벤트를 기록한다. 이 이벤트는 이후 그룹매칭 퍼널의 각 단계별 수치로 집계된다. 따라서 운영자는 단순히 몇 개의 그룹방이 만들어졌는지가 아니라, 어느 단계에서 사용자가 멈추는지 확인할 수 있다.",
            "사용자 활동 데이터가 통계로 이어지는 흐름도 중요하다. 사용자가 앱에 진입하거나 주요 화면을 볼 때 클라이언트 이벤트가 저장되고, 서버는 매칭 요청, 요청 수락, 그룹매칭 완료, 티켓 지급, 결제 완료와 같은 결과 이벤트를 별도로 저장한다. 두 종류의 이벤트가 결합되면 사용자의 행동 의도와 서버의 처리 결과를 함께 분석할 수 있다. 이는 기능 개선의 우선순위를 정하는 데 유용하다.",
            "제공된 시퀀스 도식은 이러한 연계가 순차적으로 발생함을 보여준다. 개인 매칭 시퀀스에서는 추천 후보 조회, 요청 생성, 상대방 알림, 요청 수락, 개인 대화 공간 생성, 수락 알림이 이어진다. 그룹매칭 시퀀스에서는 팀방 생성, 팀원 참여, 준비 상태 변경, 대기열 진입, 상대 팀 탐색, 최종 그룹 대화 공간 생성, 이동 알림이 이어진다. 각 단계는 단일 기능이 아니라 여러 도메인이 함께 처리하는 서비스 흐름이다.",
        ],
        None,
    ),
    (
        "9",
        "백엔드 설계상 주요 고려사항",
        [
            "첫째, 책임 분리이다. AirConnect는 인증, 사용자, 매칭, 채팅, 알림, 그룹매칭, 통계, 운영 관리가 서로 다른 책임을 갖도록 나뉘어 있다. 기능 간 연계가 많더라도 각 도메인은 자신의 상태와 규칙을 중심으로 처리하고, 다른 기능과의 연결은 필요한 시점에 명확한 요청으로 수행한다. 이 구조는 유지보수성과 테스트 가능성을 높인다.",
            "둘째, 데이터 정합성이다. 매칭이 수락되었는데 채팅방이 없거나 알림 발송 대기 기록이 원본 알림을 찾지 못하거나 티켓 잔액이 사용 이력과 맞지 않는 상황은 서비스 신뢰도를 떨어뜨린다. 이를 방지하기 위해 주요 상태 변경은 트랜잭션 안에서 처리되고, 운영 기능에서는 정합성 점검 지표를 제공한다. 데이터 정합성은 기능 품질뿐 아니라 운영 신뢰성의 핵심이다.",
            "셋째, 트랜잭션 처리이다. 채팅방 생성과 참여자 등록, 메시지 저장과 채팅방 최신 상태 갱신, 그룹매칭 완료와 최종 그룹방 생성, 티켓 차감, 결제 검증 후 지급과 같은 흐름은 원자적 처리가 필요하다. 특히 티켓과 결제, 광고 보상은 중복 지급과 동시 차감 문제가 발생할 수 있으므로 사용자 잔액 또는 주문 상태를 잠근 뒤 처리하는 방식이 적용된다.",
            "넷째, 예외 처리이다. AirConnect의 기능은 사용자 상태, 프로필 완성도, 차단 관계, 티켓 잔액, 팀방 상태, 준비 여부, 대기열 상태, 알림 설정과 같은 조건을 지속적으로 검증한다. 예외 처리는 단순 오류 응답이 아니라 서비스 규칙을 보호하는 장치이다. 예를 들어 차단 관계에서는 매칭과 채팅이 제한되고, 준비가 완료되지 않은 팀은 대기열에 들어갈 수 없으며, 만료된 광고 보상 세션은 지급되지 않는다.",
            "다섯째, 확장성과 유지보수성이다. 도메인별 모델과 계층 구조가 분리되어 있어 새로운 매칭 조건, 새로운 알림 유형, 새로운 통계 지표, 새로운 결제 상품을 추가할 수 있다. 알림은 원본 기록과 발송 대기 구조가 분리되어 있어 푸시 제공자가 바뀌어도 핵심 알림 기록 구조를 유지할 수 있다. 그룹매칭은 팀 크기별 대기열 구조를 기반으로 추가 조건을 확장할 수 있다.",
            "여섯째, 성능 고려이다. 채팅 메시지는 커서 기반으로 조회되고, 채팅방 목록은 최신 메시지와 미읽음 수를 함께 계산한다. 알림 목록도 커서 기반으로 제공되어 많은 알림을 가진 사용자에게도 안정적으로 동작한다. 통계 기능은 필요한 집계만 읽기 전용 흐름으로 수행하여 사용자 요청 처리와 분리된다. 실시간 통신은 필요한 이벤트만 구독자에게 전달함으로써 불필요한 부하를 줄인다.",
            "일곱째, 동시성 및 실시간 처리 고려이다. 그룹매칭 대기열은 여러 작업자가 동시에 같은 팀을 처리하지 않도록 잠금과 재시도를 사용한다. 채팅 읽음 처리는 사용자가 여러 기기 또는 여러 구독 경로로 접속해도 마지막 읽음 기준을 유지한다. 알림 발송 대기열은 발송 가능한 항목을 점유하고, 실패 시 재시도하거나 오래 묶인 항목을 복구한다. 이러한 설계는 실제 운영 환경의 불안정성을 전제로 한 것이다.",
        ],
        None,
    ),
    (
        "10",
        "구현 결과 및 기대 효과",
        [
            "구현 결과 사용자는 매칭 이후 자연스럽게 대화로 이어지는 경험을 얻는다. 개인 매칭은 수락 후 바로 개인 대화 공간으로 이어지고, 그룹매칭은 팀방에서 준비를 마친 뒤 최종 그룹 대화 공간으로 이동한다. 채팅방 목록과 미읽음 수, 읽음 처리, 실시간 메시지 전달이 함께 제공되므로 사용자는 현재 대화 상태를 명확히 파악할 수 있다.",
            "알림 기능은 사용자의 참여 지속성을 높인다. 매칭 요청, 팀원 합류, 준비 상태 변경, 최종 그룹매칭 성사, 새 메시지, 운영 공지가 알림으로 전달되기 때문에 사용자는 서비스 흐름에서 이탈하더라도 중요한 상태 변화를 다시 확인할 수 있다. 사용자별 알림 설정과 방해금지 시간대는 과도한 푸시를 줄이고, 발송 대기 구조는 외부 서비스 장애에도 회복 가능한 운영 구조를 제공한다.",
            "그룹매칭 기능은 AirConnect의 차별적인 서비스 가치를 만든다. 팀 구성과 준비 확인을 서버가 엄격하게 관리하고, 대기열과 조건 기반 매칭을 통해 상대 팀을 결정하며, 최종 대화 공간까지 자동으로 연결한다. 사용자는 복잡한 팀 단위 약속 과정을 서비스 내부에서 처리할 수 있고, 운영자는 각 단계의 상태와 결과를 데이터로 확인할 수 있다.",
            "통계 기능은 서비스 운영 측면의 이점을 제공한다. 전체 사용자 규모, 당일 활동, 성별 구성, 매칭 성과, 학과별 요청 순위, 매칭 퍼널, 그룹매칭 퍼널, 알림 발송 상태, 신고 처리 현황을 통해 운영자는 기능 개선 방향을 판단할 수 있다. 교수자 평가 관점에서도 단순 구현이 아니라 서비스 운영까지 고려한 백엔드 설계라는 점이 드러난다.",
            "기술적 완성도 측면에서는 관계형 데이터베이스 기반 상태 관리, 트랜잭션 처리, 실시간 통신, 비동기 알림 발송, 대기열 복구, 티켓 사용 이력 관리, 운영 통계 집계가 결합되어 있다. 이는 수업 과제 수준의 단순 기능 목록을 넘어 실제 서비스 운영을 염두에 둔 구조적 구현이라는 의미를 가진다.",
        ],
        None,
    ),
    (
        "11",
        "결론",
        [
            "AirConnect 백엔드의 채팅, 알림, 그룹매칭, 통계 기능은 서로 독립적인 기능처럼 보이지만 실제로는 하나의 서비스 흐름을 구성한다. 그룹매칭은 사용자를 새로운 관계로 연결하고, 채팅은 그 관계가 실제 대화로 이어지는 공간을 제공한다. 알림은 사용자가 상태 변화를 놓치지 않도록 연결하며, 통계는 이러한 활동을 운영 가능한 정보로 전환한다.",
            "네 기능의 핵심 성과는 상태 전이의 명확성, 데이터 영속성, 실시간성, 비동기 처리, 운영 분석 가능성으로 요약할 수 있다. 채팅 기능은 메시지 저장과 읽음 처리를 안정적으로 관리하고, 알림 기능은 인앱 기록과 외부 푸시 발송을 분리한다. 그룹매칭 기능은 팀 구성부터 최종 그룹 대화방 생성까지 복잡한 흐름을 단계적으로 제어하며, 통계 기능은 서비스 현황과 사용자 행동을 정량적으로 해석할 수 있게 한다.",
            "AirConnect 백엔드에서 이 기능들이 가지는 의미는 단순한 편의 기능을 넘어선다. 이들은 사용자 경험, 서비스 신뢰성, 운영 가능성, 확장성을 동시에 지탱하는 핵심 기반이다. 향후에는 관심사 기반 그룹매칭 조건 확장, 궁합 결과와 매칭 추천의 결합, 알림 개인화, 통계 기반 운영 자동화, 안전 기능의 고도화로 발전할 수 있다. 이러한 확장 가능성은 현재 구조가 도메인별 책임 분리와 데이터 흐름 중심으로 설계되어 있기 때문에 가능하다.",
            "따라서 AirConnect의 백엔드 구현은 기능별 요구사항을 충족하는 데 그치지 않고, 사용자 연결 서비스가 갖추어야 할 신뢰성, 일관성, 실시간성, 운영성을 함께 고려한 설계 결과물로 평가할 수 있다.",
        ],
        None,
    ),
]


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = add_paragraph(doc, "AirConnect 백엔드 기능 설계 및 구현 보고서", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.runs[0], size=24, bold=True, color="0B2545")
    p.paragraph_format.space_after = Pt(18)

    p = add_paragraph(doc, "주제: 인증, 회원, 개인 매칭, 궁합, 결제, 광고 보상, 티켓, 운영 기능 중심", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.runs[0], size=13.5, bold=True, color="1F4D78")
    p.paragraph_format.space_after = Pt(36)

    for label in ["이름 : 정재민", "소속 : ______________________________", "학번 : ______________________________", "제출일 : ____________________________"]:
        p = add_paragraph(doc, label, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_font(p.runs[0], size=12)
        p.paragraph_format.space_after = Pt(10)

    for _ in range(5):
        doc.add_paragraph()
    p = add_paragraph(doc, "제출용 백엔드 설계 보고서", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.runs[0], size=11, color="555555")
    page_break(doc)


def add_toc(doc):
    add_heading(doc, "", "목차", 1)
    items = [
        "1. 프로젝트 개요",
        "2. 전체 백엔드 설계 개요",
        "3. 데이터 모델 및 ERD 분석",
        "4. 인증 및 회원 기능 설계",
        "5. 개인 매칭 및 후보 추천 기능 설계",
        "6. 궁합 분석 기능 설계",
        "7. 결제, 광고 보상, 티켓 기능 설계",
        "8. 운영 관리, 공지, 점검 기능 설계",
        "9. 안전 관리 기능 설계",
        "10. 주요 처리 흐름도",
        "11. 백엔드 설계상 주요 고려사항",
        "12. 구현 결과 및 결론",
    ]
    for item in items:
        p = add_paragraph(doc, item)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(5)
    page_break(doc)


def create_scope_erd():
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    path = DIAGRAM_DIR / "project_scope_erd.png"
    width, height = 2300, 1500
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font = load_diagram_font(42, bold=True)
    box_title_font = load_diagram_font(25, bold=True)
    item_font = load_diagram_font(21)
    rel_font = load_diagram_font(19)

    title = "작성 범위 기준 핵심 ERD"
    title_width, _ = text_size(draw, title, title_font)
    draw.text(((width - title_width) / 2, 30), title, font=title_font, fill="#1F4D78")
    subtitle = "인증, 회원, 개인 매칭, 궁합, 결제, 광고 보상, 티켓, 운영 기능 중심 데이터 관계"
    subtitle_width, _ = text_size(draw, subtitle, item_font)
    draw.text(((width - subtitle_width) / 2, 86), subtitle, font=item_font, fill="#667085")

    palette = {
        "user": ("#EAF4FF", "#2E74B5"),
        "match": ("#F1F8FF", "#5B8DEF"),
        "compat": ("#F3F0FF", "#6F4CEB"),
        "pay": ("#FFF4F4", "#C2410C"),
        "ops": ("#F5F5F5", "#667085"),
        "safe": ("#FFF7E6", "#9A6700"),
    }
    row_labels = [
        ("회원과 인증", 155, "#EAF4FF"),
        ("개인 매칭과 궁합", 405, "#F3F0FF"),
        ("결제, 광고 보상, 티켓", 655, "#FFF4F4"),
        ("운영 관리와 서비스 제어", 905, "#F5F5F5"),
        ("안전 관리", 1155, "#FFF7E6"),
    ]
    for label, y, fill in row_labels:
        draw.rounded_rectangle((360, y - 42, 2220, y + 170), radius=18, fill=fill, outline="#E4E7EC", width=2)
        draw.text((378, y - 34), label, font=item_font, fill="#344054")

    boxes = {
        "사용자 정보": (70, 650, 260, 140, "user", ["계정 상태", "역할과 보유 티켓"]),
        "프로필 정보": (440, 155, 290, 122, "user", ["닉네임, 학교, 학과", "관심사와 취향"]),
        "외부 인증 기록": (800, 155, 290, 122, "user", ["카카오와 애플 인증", "기기 연결 정보"]),
        "학교 검증 정보": (1160, 155, 290, 122, "user", ["학교 이메일 확인", "검증 완료 시각"]),
        "탈퇴 처리 기록": (1520, 155, 290, 122, "user", ["계정 비활성화", "재가입 제한 기준"]),
        "후보 노출 기록": (440, 405, 290, 122, "match", ["후보 표시 이력", "중복 추천 방지"]),
        "개인 매칭 요청": (800, 405, 290, 122, "match", ["요청, 수락, 거절", "중복 연결 방지"]),
        "궁합 분석 결과": (1160, 405, 290, 122, "compat", ["8개 항목 점수", "요약 문장과 재사용"]),
        "티켓 잔액": (440, 655, 290, 122, "pay", ["사용자별 보유 수량", "차감 가능 기준"]),
        "티켓 사용 이력": (800, 655, 290, 122, "pay", ["지급, 차감 기록", "중복 처리 방지"]),
        "결제 검증 기록": (1160, 655, 290, 122, "pay", ["스토어 검증 결과", "계정 식별값 확인"]),
        "광고 보상 기록": (1520, 655, 290, 122, "pay", ["보상 세션", "서명 검증 결과"]),
        "운영 계정": (440, 905, 290, 122, "ops", ["운영자 권한", "접근 범위 구분"]),
        "운영 작업 기록": (800, 905, 290, 122, "ops", ["조회와 변경 이력", "감사 가능성 확보"]),
        "공지 정보": (1160, 905, 290, 122, "ops", ["노출 기간", "우선순위"]),
        "점검 설정": (1520, 905, 290, 122, "ops", ["점검 모드", "버전 제어"]),
        "신고 기록": (440, 1155, 290, 122, "safe", ["신고 대상과 사유", "처리 상태"]),
        "차단 목록": (800, 1155, 290, 122, "safe", ["사용자 간 제한", "재요청 차단"]),
        "제재 기록": (1160, 1155, 290, 122, "safe", ["블랙리스트", "서비스 이용 제한"]),
    }

    box_rects = {}
    for name, (x, y, w, h, kind, items) in boxes.items():
        fill, outline = palette[kind]
        draw_entity_box(draw, x, y, w, h, name, items, fill, outline, box_title_font, item_font)
        box_rects[name] = (x, y, w, h)

    bus_x = 370
    user_x, user_y, user_w, user_h = box_rects["사용자 정보"]
    draw.line((user_x + user_w, user_y + user_h / 2, bus_x, user_y + user_h / 2), fill="#667085", width=4)
    draw.line((bus_x, 215, bus_x, 1220), fill="#667085", width=4)

    def connect_user(target, label):
        x, y, w, h = box_rects[target]
        mid_y = y + h / 2
        draw.line((bus_x, mid_y, x, mid_y), fill="#667085", width=3)
        tw, th = text_size(draw, label, rel_font)
        draw.rounded_rectangle((bus_x + 14, mid_y - th - 16, bus_x + tw + 34, mid_y - 4), radius=8, fill="#FFFFFF")
        draw.text((bus_x + 22, mid_y - th - 13), label, font=rel_font, fill="#475467")

    for target, label in [
        ("프로필 정보", "1:1"),
        ("외부 인증 기록", "1:N"),
        ("학교 검증 정보", "1:N"),
        ("후보 노출 기록", "1:N"),
        ("티켓 잔액", "1:1"),
        ("운영 계정", "조건부"),
        ("신고 기록", "1:N"),
    ]:
        connect_user(target, label)

    def connect_row(start, end, label):
        x1, y1, w1, h1 = box_rects[start]
        x2, y2, w2, h2 = box_rects[end]
        sx, sy = x1 + w1, y1 + h1 / 2
        ex, ey = x2, y2 + h2 / 2
        draw.line((sx, sy, ex, ey), fill="#667085", width=3)
        draw.ellipse((ex - 5, ey - 5, ex + 5, ey + 5), fill="#667085")
        tw, th = text_size(draw, label, rel_font)
        tx, ty = (sx + ex) / 2, (sy + ey) / 2
        draw.rounded_rectangle((tx - tw / 2 - 8, ty - th - 12, tx + tw / 2 + 8, ty - 2), radius=8, fill="#FFFFFF")
        draw.text((tx - tw / 2, ty - th - 9), label, font=rel_font, fill="#475467")

    for start, end, label in [
        ("외부 인증 기록", "학교 검증 정보", "보완"),
        ("학교 검증 정보", "탈퇴 처리 기록", "상태"),
        ("후보 노출 기록", "개인 매칭 요청", "근거"),
        ("개인 매칭 요청", "궁합 분석 결과", "선택"),
        ("티켓 잔액", "티켓 사용 이력", "1:N"),
        ("티켓 사용 이력", "결제 검증 기록", "참조"),
        ("티켓 사용 이력", "광고 보상 기록", "참조"),
        ("운영 계정", "운영 작업 기록", "1:N"),
        ("운영 작업 기록", "공지 정보", "관리"),
        ("공지 정보", "점검 설정", "서비스 제어"),
        ("신고 기록", "차단 목록", "처리"),
        ("차단 목록", "제재 기록", "확대"),
    ]:
        connect_row(start, end, label)

    legend = "관계 표기: 1:1 단일 관계, 1:N 다중 이력, 조건부는 권한 또는 상태에 따른 연결"
    legend_width, _ = text_size(draw, legend, item_font)
    draw.rounded_rectangle((width - legend_width - 110, 1390, width - 60, 1432), radius=12, fill="#F2F4F7")
    draw.text((width - legend_width - 86, 1398), legend, font=item_font, fill="#475467")
    image.save(path, quality=95)
    return path


def add_project_erd(doc):
    p = add_paragraph(doc, "그림 1. 작성 범위 기준 핵심 ERD", keep_with_next=True)
    set_font(p.runs[0], size=10, bold=True, color="1F4D78")
    picture_paragraph = doc.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(str(create_scope_erd()), width=Inches(6.65))
    picture_paragraph.paragraph_format.space_after = Pt(8)


def ensure_sequence_diagrams():
    specs = [
        (
            "sequence_auth_scope.png",
            "인증 및 회원 등록 시퀀스 흐름도",
            ["사용자", "모바일 앱", "외부 인증 제공자", "요청 접수 계층", "인증 처리 계층", "데이터 저장소"],
            [
                (0, 1, "소셜 인증 시작", False),
                (1, 2, "외부 인증 진행", False),
                (2, 1, "인증 결과 반환", True),
                (1, 3, "인증 정보 전달", False),
                (3, 4, "토큰 검증과 사용자 조회", False),
                (4, 5, "기존 계정과 상태 확인", False),
                (5, 4, "계정 데이터 반환", True),
                (4, 5, "토큰 발급 정보 저장", False),
                (4, 3, "인증 처리 결과 반환", True),
                (3, 1, "서비스 진입 결과 전달", True),
            ],
            [(5, 4, "계정 상태, 탈퇴 여부, 기기 연결 확인"), (8, 5, "갱신 가능한 인증 상태 보존")],
            [(4, 8, "검증")],
            "2E74B5",
            "그림 2. 인증 및 회원 등록 시퀀스 흐름도",
        ),
        (
            "sequence_matching_scope.png",
            "개인 매칭 후보 추천 시퀀스 흐름도",
            ["사용자", "모바일 앱", "요청 접수 계층", "추천 처리 계층", "데이터 저장소", "티켓 처리 계층"],
            [
                (0, 1, "후보 찾기 진입", False),
                (1, 2, "후보 추천 요청", False),
                (2, 3, "추천 처리 호출", False),
                (3, 4, "프로필 조건과 차단 목록 조회", False),
                (4, 3, "후보 기초 데이터 반환", True),
                (3, 4, "노출 이력과 요청 상태 확인", False),
                (3, 5, "티켓 사용 가능 여부 확인", False),
                (5, 4, "티켓 사용 이력 저장", False),
                (3, 2, "최종 후보 목록 반환", True),
                (2, 1, "화면에 후보 표시", True),
            ],
            [(4, 4, "본인 제외, 조건 일치, 중복 추천 방지"), (7, 5, "필요한 경우에만 티켓 차감")],
            [(3, 6, "필터링"), (7, 8, "티켓")],
            "6F4CEB",
            "그림 3. 개인 매칭 후보 추천 시퀀스 흐름도",
        ),
        (
            "sequence_ad_scope.png",
            "광고 보상 시퀀스 흐름도",
            ["사용자", "모바일 앱", "광고 플랫폼", "요청 접수 계층", "보상 처리 계층", "데이터 저장소"],
            [
                (0, 1, "광고 시청 진입", False),
                (1, 3, "보상 세션 발급 요청", False),
                (3, 4, "세션 생성 처리 호출", False),
                (4, 5, "보상 세션 저장", False),
                (5, 4, "세션 정보 반환", True),
                (1, 2, "광고 시청 완료", False),
                (2, 3, "완료 통지 전달", False),
                (3, 4, "전자 서명과 전송 정보 검증", False),
                (4, 5, "세션 상태와 중복 보상 조회", False),
                (4, 5, "티켓 사용 이력 저장", False),
                (4, 5, "잔액 반영", False),
                (4, 3, "보상 처리 결과 반환", True),
                (3, 1, "결과 동기화", True),
            ],
            [(8, 4, "서명 검증, 만료 여부, 중복 지급 확인"), (10, 5, "사용 이력 기준으로 잔액 변화 추적")],
            [(2, 5, "세션"), (7, 11, "보상")],
            "00856F",
            "그림 4. 광고 보상 시퀀스 흐름도",
        ),
        (
            "sequence_payment_scope.png",
            "결제 및 티켓 지급 시퀀스 흐름도",
            ["사용자", "모바일 앱", "스토어 검증 서버", "요청 접수 계층", "결제 처리 계층", "데이터 저장소"],
            [
                (0, 1, "상품 결제 완료", False),
                (1, 3, "결제 정보 전달", False),
                (3, 4, "결제 검증 처리 호출", False),
                (4, 2, "스토어 서버 검증 요청", False),
                (2, 4, "검증 결과 반환", True),
                (4, 5, "주문 상태 잠금", False),
                (4, 5, "계정 식별과 중복 지급 확인", False),
                (4, 5, "상품 정책에 따른 지급 결정", False),
                (4, 5, "티켓 사용 이력 저장", False),
                (4, 5, "잔액 반영", False),
                (4, 3, "결제 처리 결과 반환", True),
                (3, 1, "구매 결과와 잔액 갱신", True),
            ],
            [(4, 2, "구매 정보의 유효성과 환경 확인"), (7, 5, "동일 결제의 반복 지급 차단")],
            [(3, 8, "검증"), (9, 10, "지급")],
            "9A6700",
            "그림 5. 결제 및 티켓 지급 시퀀스 흐름도",
        ),
    ]
    return [
        (caption, create_sequence_diagram(filename, title, lanes, steps, notes, groups, accent))
        for filename, title, lanes, steps, notes, groups, accent, caption in specs
    ]


def add_flow_diagrams(doc):
    add_heading(doc, "", "주요 처리 흐름도", 2)
    add_paragraph(
        doc,
        "다음 도식은 작성 범위에 포함되는 기능만을 대상으로 구성하였다. 각 도식은 실제 코드 식별자와 구체적인 요청 경로를 사용하지 않고, 백엔드 계층과 데이터 흐름을 개념명으로 표현한다.",
    )
    for caption, path in ensure_sequence_diagrams():
        p = add_paragraph(doc, caption, keep_with_next=True)
        set_font(p.runs[0], size=10, bold=True, color="1F4D78")
        picture_paragraph = doc.add_paragraph()
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.add_run().add_picture(str(path), width=Inches(6.65))
        picture_paragraph.paragraph_format.space_after = Pt(8)


SECTIONS = [
    (
        "1",
        "프로젝트 개요",
        [
            "AirConnect는 대학생이 신뢰 가능한 인증과 프로필 정보를 바탕으로 새로운 연결을 탐색할 수 있도록 설계된 캠퍼스 기반 서비스이다. 본 보고서는 작성자가 구현한 범위에 해당하는 인증, 회원, 개인 매칭, 궁합, 결제, 광고 보상, 티켓, 운영 관리, 공지, 점검, 안전 기능을 중심으로 설명한다. 보고서의 목적은 단순한 기능 나열이 아니라 백엔드 설계 방향, 데이터 흐름, 상태 관리 방식, 기술적 의사결정을 제출용 문서 형식으로 정리하는 데 있다.",
            "서비스 관점에서 AirConnect가 해결하려는 문제는 사용자가 신뢰 가능한 상대를 탐색하고 서비스 안에서 필요한 조건과 비용을 명확히 인지하며 운영 정책의 보호를 받도록 하는 것이다. 이를 위해 소셜 인증과 학교 검증으로 사용자 신뢰 기반을 마련하고, 프로필과 취향 정보를 바탕으로 개인 추천을 제공하며, 궁합 분석으로 후보 이해를 돕는다. 또한 결제와 광고 보상, 티켓 사용 이력은 서비스 이용 비용과 보상의 근거를 투명하게 남긴다.",
            "전체 백엔드 구조는 기능별 책임 분리와 상태 기반 처리를 지향한다. 인증 흐름은 사용자 신원을 확인하고 서비스 진입 가능성을 판단하며, 회원 기능은 프로필과 관심사 데이터를 관리한다. 개인 매칭 기능은 후보 노출, 요청 상태, 중복 연결 방지를 담당하고, 궁합 기능은 외부 인공지능 모델을 활용해 후보 정보를 해석 가능한 결과로 변환한다. 결제와 광고 보상은 티켓 잔액과 사용 이력을 통해 중복 지급과 부정 이용을 방지한다.",
            "운영 관리, 공지, 점검, 안전 기능은 실제 서비스 운영을 위한 기반이다. 운영자는 사용자 상태와 신고 처리를 확인하고, 공지와 점검 설정을 통해 서비스 노출과 이용 가능 상태를 조정할 수 있다. 안전 기능은 신고 접수, 차단 목록, 제재 기록을 통해 사용자 보호와 서비스 신뢰성을 유지한다.",
        ],
        None,
    ),
    (
        "2",
        "전체 백엔드 설계 개요",
        [
            "AirConnect 백엔드는 계층형 구조를 기반으로 한다. 클라이언트 요청은 인증 및 권한 검사를 거친 뒤 도메인별 처리 흐름으로 전달되고, 각 도메인은 필요한 데이터 검증과 상태 변경을 수행한다. 상태 변경이 필요한 기능은 트랜잭션 단위로 처리되어 데이터 정합성을 유지하고, 조회 중심 기능은 읽기 전용 흐름으로 분리되어 불필요한 변경을 방지한다.",
            "도메인 중심 설계의 핵심은 서비스의 주요 상태를 기능별 개념으로 분리한 점이다. 사용자 정보, 프로필 정보, 외부 인증 기록, 학교 검증 정보, 후보 노출 기록, 개인 매칭 요청, 궁합 분석 결과, 결제 검증 기록, 광고 보상 기록, 티켓 잔액, 티켓 사용 이력, 운영 작업 기록, 공지 정보, 점검 설정, 신고 기록, 차단 목록, 제재 기록이 각각 독립된 의미를 갖는다.",
            "데이터베이스 기반 상태 관리는 중요한 설계 선택이다. 인증 토큰의 발급과 갱신, 프로필 등록과 수정, 후보 추천 이력, 요청 상태, 궁합 결과 재사용, 결제 검증 상태, 보상 세션 상태, 티켓 지급과 차감, 공지 노출 기간, 점검 모드, 신고 처리 상태는 모두 명확한 상태값과 이력으로 관리된다. 이를 통해 사용자가 같은 요청을 반복하거나 네트워크가 불안정한 상황에서도 서버가 최종 상태를 기준으로 일관된 응답을 제공할 수 있다.",
            "외부 서비스와 연결되는 기능은 실패 가능성을 전제로 설계된다. 소셜 인증 제공자, 스토어 검증 서버, 광고 플랫폼, 인공지능 모델은 모두 서버 바깥의 시스템이므로 백엔드는 검증 결과를 저장하고, 중복 처리 기준을 마련하며, 실패 시 사용자 상태가 잘못 변경되지 않도록 경계를 둔다. 이러한 구조는 외부 호출이 많은 모바일 서비스에서 특히 중요하다.",
        ],
        {
            "headers": ["설계 축", "적용 방식", "의미"],
            "rows": [
                ["책임 분리", "기능별 도메인과 계층형 처리", "변경 영향 범위 축소"],
                ["상태 관리", "관계형 데이터 기반 상태 전이", "반복 요청과 예외 상황 대응"],
                ["검증 구조", "외부 결과와 내부 상태의 분리", "보안성과 운영 안정성 확보"],
                ["티켓 관리", "잔액과 사용 이력 동시 관리", "중복 지급과 동시 차감 방지"],
                ["운영 제어", "공지, 점검, 제재 상태 관리", "서비스 운영 가능성 확보"],
            ],
            "widths": [1800, 3300, 4260],
        },
    ),
    (
        "3",
        "데이터 모델 및 ERD 분석",
        [
            "제공된 ERD 도식과 프로젝트 구현을 함께 검토하면 작성 범위의 데이터 모델은 회원과 인증, 개인 매칭과 궁합, 결제와 광고 보상, 티켓, 운영 관리, 공지와 점검, 안전 관리로 나뉜다. 아래 핵심 ERD는 실제 프로젝트 스캔 결과를 바탕으로 주요 영속 모델을 코드 식별자가 아닌 기능 개념명으로 재구성한 것이다.",
            "회원과 인증 영역에서는 사용자 정보가 프로필 정보, 외부 인증 기록, 학교 검증 정보, 탈퇴 처리 기록과 연결된다. 이 관계는 인증과 사용자 상태가 이후 기능의 선행 조건임을 의미한다. 프로필 정보는 추천 조건과 궁합 분석의 입력이 되며, 학교 검증 정보는 서비스 신뢰성을 높이는 기준으로 작동한다.",
            "개인 매칭과 궁합 영역에서는 후보 노출 기록, 개인 매칭 요청, 궁합 분석 결과가 연결된다. 후보 노출 기록은 동일 후보가 반복적으로 노출되는 문제를 줄이고, 개인 매칭 요청은 요청 상태와 중복 연결 방지를 담당한다. 궁합 분석 결과는 사용자의 프로필과 취향 데이터를 해석해 점수와 요약 문장으로 저장되며, 같은 조건의 반복 계산을 줄이기 위해 재사용 가능한 결과로 관리된다.",
            "결제, 광고 보상, 티켓 영역에서는 티켓 잔액과 티켓 사용 이력이 중심이 된다. 결제 검증 기록과 광고 보상 기록은 모두 티켓 사용 이력과 연결되어 잔액 변화의 원인을 남긴다. 이 구조는 티켓을 단순 숫자로만 저장하지 않고, 언제 어떤 이유로 지급 또는 차감되었는지 확인할 수 있게 한다.",
            "운영 관리와 안전 영역에서는 운영 계정, 운영 작업 기록, 공지 정보, 점검 설정, 신고 기록, 차단 목록, 제재 기록이 연결된다. 운영자는 공지 노출과 점검 상태를 관리하고, 신고와 차단, 제재 데이터를 통해 서비스 안전성을 유지한다. 결과적으로 ERD는 작성 범위의 기능들이 사용자 신뢰, 비용 처리, 운영 제어, 안전 관리라는 네 축으로 구성되어 있음을 보여준다.",
        ],
        {
            "headers": ["개념 영역", "주요 관계", "설계 의미"],
            "rows": [
                ["회원과 인증", "사용자 정보와 프로필, 외부 인증, 학교 검증", "서비스 진입 가능성 판단"],
                ["개인 매칭", "후보 노출 기록과 요청 상태", "중복 추천과 중복 연결 방지"],
                ["궁합", "프로필 기반 분석 결과와 재사용", "외부 모델 호출 비용 절감"],
                ["결제와 보상", "검증 기록과 티켓 사용 이력", "중복 지급과 잔액 오류 방지"],
                ["운영과 안전", "공지, 점검, 신고, 차단, 제재 기록", "운영 가능성과 사용자 보호 확보"],
            ],
            "widths": [1900, 3500, 3960],
        },
    ),
    (
        "4",
        "인증 및 회원 기능 설계",
        [
            "인증 기능의 목적은 사용자가 신뢰 가능한 방식으로 서비스에 진입하도록 하는 것이다. AirConnect는 소셜 인증을 통해 외부 인증 제공자의 검증 결과를 확인하고, 서버는 해당 결과를 내부 사용자 상태와 연결한다. 이때 외부 인증 결과를 그대로 신뢰하는 것이 아니라 계정 상태, 탈퇴 여부, 기기 연결 상태, 토큰 갱신 가능성을 함께 검토한다.",
            "토큰 발급과 갱신은 사용자 경험과 보안을 함께 고려한다. 사용자가 로그인하면 서버는 서비스 접근에 필요한 인증 상태를 발급하고, 갱신 흐름에서는 기존 인증 상태의 유효성과 사용자 상태를 다시 확인한다. 탈퇴 처리된 계정이나 제한된 계정은 정상 사용자와 동일하게 다루지 않으며, 서비스 진입 가능성을 별도로 판단한다.",
            "회원 기능은 프로필 등록, 수정, 관심사와 취향 데이터 관리를 담당한다. 프로필은 단순 표시 정보가 아니라 후보 추천과 궁합 분석에 활용되는 핵심 입력값이다. 따라서 프로필 완성도, 학교 정보, 학과 정보, 성별, 관심사, 선호 조건은 사용자 경험과 추천 품질에 직접적인 영향을 준다.",
            "차단 목록은 회원 데이터와 안전 기능을 연결한다. 사용자가 특정 상대와의 연결을 원하지 않는 경우 해당 관계는 후보 추천과 요청 처리에서 제외 기준으로 사용된다. 이는 단순 편의 기능이 아니라 사용자가 서비스 안에서 통제권을 갖도록 하는 안전 장치이다.",
        ],
        None,
    ),
    (
        "5",
        "개인 매칭 및 후보 추천 기능 설계",
        [
            "개인 매칭 기능의 목적은 사용자 프로필과 선호 조건을 바탕으로 적절한 후보를 추출하고, 반복 노출과 중복 요청을 줄이는 것이다. 후보 추천은 무작위 목록 제공이 아니라 사용자 상태, 프로필 조건, 차단 목록, 과거 노출 이력, 기존 요청 상태를 함께 검토하는 흐름으로 설계된다.",
            "후보 추출 단계에서는 먼저 본인을 제외하고 기본 조건에 맞는 후보군을 구성한다. 이후 성별 조건, 학교와 학과 정보, 관심사, 차단 관계, 이미 본 후보 여부, 이미 요청한 관계 여부를 확인한다. 이러한 필터링은 데이터베이스 조회 단계와 응용 계층의 추가 검증으로 나뉘어 수행된다.",
            "요청 상태 관리는 중복 연결 방지를 위해 중요하다. 사용자가 특정 상대에게 요청을 보내면 요청 기록은 대기, 수락, 거절과 같은 상태로 관리된다. 같은 두 사용자 사이의 관계는 하나의 기준으로 정규화되어 중복 요청이 만들어지지 않도록 설계된다. 거절 이후 재요청이 필요한 경우에도 이전 상태와 요청자를 기준으로 일관된 흐름을 유지한다.",
            "티켓 차감은 후보 추천 또는 연결 요청과 결합될 수 있다. 이때 서버는 사용자의 보유 티켓을 확인하고, 실제 차감이 필요한 조건인지 판단한 뒤 티켓 사용 이력을 남긴다. 티켓 차감과 추천 결과 생성은 같은 흐름 안에서 일관되게 처리되어 사용자가 결과를 받았지만 비용 처리가 실패하거나 반대로 비용만 차감되는 상황을 방지한다.",
        ],
        None,
    ),
    (
        "6",
        "궁합 분석 기능 설계",
        [
            "궁합 기능은 사용자의 프로필과 취향 데이터를 바탕으로 상대 후보와의 적합도를 설명 가능한 형태로 제공하는 기능이다. 단순 점수 하나를 제공하는 것이 아니라 여러 항목을 나누어 점수화하고, 사용자가 이해할 수 있는 요약 문장을 함께 생성한다. 이를 통해 추천 결과가 왜 의미 있는지 설명하는 보조 정보를 제공한다.",
            "외부 인공지능 모델을 활용하는 기능은 입력 데이터의 정리와 결과 재사용이 중요하다. 서버는 분석에 필요한 프로필 요소를 구성하고, 외부 모델 호출 결과를 항목별 점수와 요약 문장으로 정리한다. 동일한 조건에서 반복적으로 같은 계산이 발생하지 않도록 결과를 저장해 재사용할 수 있게 한다.",
            "결과 캐싱은 비용과 응답 시간을 줄이는 설계 선택이다. 궁합 분석은 일반 조회보다 비용이 큰 기능이므로 요청마다 외부 모델을 호출하면 성능과 비용 측면에서 부담이 커진다. 따라서 이미 계산된 결과가 유효한 경우 저장된 결과를 우선 활용하고, 새 분석이 필요한 경우에만 외부 호출을 수행한다.",
            "궁합 결과는 추천 기능과 분리된 독립 도메인으로 관리된다. 추천 기능이 후보 목록을 구성하는 역할을 맡는다면 궁합 기능은 후보에 대한 해석 정보를 제공한다. 이러한 분리는 향후 분석 항목 추가, 요약 문장 개선, 결과 표시 방식 변경이 발생해도 추천 후보 추출 규칙에 직접적인 영향을 주지 않도록 한다.",
        ],
        None,
    ),
    (
        "7",
        "결제, 광고 보상, 티켓 기능 설계",
        [
            "결제 기능의 목적은 사용자가 구매한 상품이 실제 스토어에서 유효한지 서버에서 검증하고, 검증된 결과에 따라 티켓을 지급하는 것이다. 모바일 환경에서는 클라이언트가 전달한 결제 정보를 그대로 신뢰할 수 없으므로 서버가 스토어 검증 서버와 통신하여 거래의 유효성, 상품 정보, 계정 식별값, 중복 지급 여부를 확인한다.",
            "광고 보상 기능은 사용자가 보상형 광고를 정상적으로 완료했을 때 티켓을 지급하는 흐름이다. 광고 플랫폼이 전달하는 완료 통지는 전자 서명과 전송 정보를 검증해야 하며, 보상 세션의 만료 여부와 중복 보상 여부를 함께 확인해야 한다. 이 과정을 거친 뒤에만 티켓이 지급된다.",
            "티켓 기능은 지급과 차감을 모두 관리한다. 티켓 잔액만 저장하면 오류 발생 시 원인을 추적하기 어렵기 때문에, 서버는 지급 또는 차감이 발생한 이유와 기준을 별도의 사용 이력으로 남긴다. 사용 이력은 결제 검증 기록, 광고 보상 기록, 추천 또는 요청 처리와 연결되어 잔액 변화의 근거를 제공한다.",
            "동시성 제어는 티켓 처리에서 특히 중요하다. 사용자가 짧은 시간에 여러 요청을 보내거나 결제 검증이 반복될 경우 같은 보상이 두 번 지급될 수 있다. 이를 방지하기 위해 주문 상태나 보상 세션 상태를 잠근 뒤 처리하고, 이미 처리된 기준값은 다시 지급하지 않도록 한다. 이러한 멱등 처리 구조는 결제와 광고 보상의 신뢰성을 높인다.",
        ],
        None,
    ),
    (
        "8",
        "운영 관리, 공지, 점검 기능 설계",
        [
            "운영 관리 기능은 서비스 운영자가 사용자 상태, 신고 처리, 티켓 조정, 공지 노출, 점검 설정을 확인하고 조정할 수 있도록 하는 기능이다. 일반 사용자 기능과 운영 기능은 권한이 명확히 분리되어야 하며, 운영 작업은 추후 확인 가능한 기록으로 남아야 한다.",
            "공지 기능은 서비스 안내를 기간과 우선순위에 따라 노출하기 위한 구조이다. 공지는 단순 텍스트 저장이 아니라 시작 시각, 종료 시각, 노출 우선순위, 활성 상태를 함께 관리한다. 이를 통해 운영자는 특정 기간에만 필요한 안내를 게시하고, 여러 공지가 동시에 존재할 때 노출 순서를 제어할 수 있다.",
            "점검 기능은 서비스 이용 가능 상태와 클라이언트 버전 정책을 제어한다. 점검 모드가 활성화되면 일반 사용자의 접근을 제한하고, 필요한 경우 최소 지원 버전이나 강제 업데이트 기준을 제공한다. 이는 장애 대응, 배포 전환, 오래된 클라이언트 차단에 필요한 운영 장치이다.",
            "운영 기능에서 중요한 점은 권한 분리와 감사 가능성이다. 운영자는 민감한 사용자 정보와 서비스 상태를 다룰 수 있기 때문에 접근 권한이 제한되어야 하고, 주요 작업은 누가 언제 어떤 변경을 했는지 확인할 수 있어야 한다. 이 구조는 운영상의 책임성과 서비스 신뢰성을 함께 높인다.",
        ],
        None,
    ),
    (
        "9",
        "안전 관리 기능 설계",
        [
            "안전 관리 기능은 신고 접수, 처리, 차단 목록, 사용자 제재를 통해 서비스 내 위험 행동을 줄이는 역할을 한다. 사용자 연결 서비스에서는 신뢰와 안전이 기능 품질만큼 중요하므로 신고와 차단, 제재 기록은 독립된 흐름으로 관리된다.",
            "신고 접수 흐름에서는 신고자와 대상자, 신고 사유, 처리 상태가 저장된다. 신고는 단순 문의가 아니라 운영 판단의 근거가 되므로 처리 전후 상태가 명확히 구분되어야 한다. 운영자는 신고 내용을 확인하고 필요에 따라 차단 또는 제재로 이어지는 조치를 수행한다.",
            "차단 목록은 사용자가 원하지 않는 상대를 서비스 흐름에서 제외하는 역할을 한다. 차단 관계는 후보 추천과 요청 처리에서 제외 기준으로 사용되며, 사용자 경험의 통제권을 보장한다. 이 기능은 서비스 안전성뿐 아니라 사용자의 심리적 부담을 낮추는 데도 기여한다.",
            "제재 기록은 반복적이거나 심각한 위반 행위에 대한 운영 조치의 결과이다. 블랙리스트 또는 이용 제한 상태는 계정 상태와 연결되어 서비스 접근 가능성을 조정한다. 안전 관리 기능은 단순 사후 처리에 그치지 않고, 이후 추천과 요청 흐름에도 영향을 주어 같은 문제가 반복되지 않도록 한다.",
        ],
        None,
    ),
    (
        "10",
        "주요 처리 흐름도",
        [
            "작성 범위의 주요 처리 흐름은 인증 및 회원 등록, 개인 매칭 후보 추천, 광고 보상, 결제 및 티켓 지급으로 정리할 수 있다. 각 흐름은 사용자 행동에서 시작하지만 서버 내부에서는 검증, 상태 조회, 중복 확인, 사용 이력 저장, 결과 반환이 단계적으로 이루어진다.",
        ],
        None,
    ),
    (
        "11",
        "백엔드 설계상 주요 고려사항",
        [
            "첫째, 책임 분리이다. 인증, 회원, 개인 매칭, 궁합, 결제, 광고 보상, 티켓, 운영, 공지, 점검, 안전 기능은 서로 다른 책임을 갖도록 분리되어 있다. 기능 간 연계가 있더라도 각 도메인은 자신의 상태와 규칙을 중심으로 처리하고, 다른 기능과의 연결은 필요한 시점에 명확한 요청으로 수행한다.",
            "둘째, 데이터 정합성이다. 결제 검증은 성공했지만 티켓이 지급되지 않거나, 광고 보상이 중복 지급되거나, 티켓 잔액이 사용 이력과 맞지 않는 상황은 서비스 신뢰도를 떨어뜨린다. 이를 방지하기 위해 주요 상태 변경은 트랜잭션 안에서 처리되고, 중복 기준값을 활용해 같은 행위가 반복 처리되지 않도록 한다.",
            "셋째, 예외 처리이다. AirConnect의 작성 범위 기능은 사용자 상태, 프로필 완성도, 학교 검증 여부, 차단 관계, 티켓 잔액, 결제 검증 결과, 광고 보상 세션 상태, 점검 모드, 제재 상태를 지속적으로 검증한다. 예외 처리는 단순 오류 응답이 아니라 서비스 규칙을 보호하는 장치이다.",
            "넷째, 확장성과 유지보수성이다. 도메인별 모델과 계층 구조가 분리되어 있어 새로운 인증 제공자, 새로운 추천 조건, 새로운 궁합 분석 항목, 새로운 결제 상품, 새로운 운영 정책을 추가할 수 있다. 특히 티켓 사용 이력 구조는 결제, 광고 보상, 운영 조정 등 여러 지급 사유를 같은 기준으로 추적할 수 있게 한다.",
            "다섯째, 보안성이다. 외부 인증, 스토어 검증, 광고 보상 검증은 모두 외부에서 전달되는 값을 다루기 때문에 서버 검증이 필수이다. 또한 운영 관리 기능은 권한 분리와 작업 기록을 통해 민감한 정보 접근을 통제해야 한다. 이러한 보안 설계는 사용자 신뢰와 운영 안정성의 기반이다.",
        ],
        None,
    ),
    (
        "12",
        "구현 결과 및 결론",
        [
            "작성 범위의 구현 결과 사용자는 신뢰 가능한 인증 절차를 거쳐 서비스에 진입하고, 프로필과 관심사 정보를 바탕으로 개인 후보를 탐색할 수 있다. 궁합 기능은 후보에 대한 이해를 돕는 분석 결과를 제공하며, 결제와 광고 보상 기능은 티켓 지급과 차감 흐름을 서버 기준으로 검증한다.",
            "운영 측면에서는 공지와 점검 설정을 통해 서비스 상태를 제어할 수 있고, 신고와 차단, 제재 기록을 통해 안전한 이용 환경을 유지할 수 있다. 티켓 사용 이력과 운영 작업 기록은 문제가 발생했을 때 원인을 추적하고 결과를 설명할 수 있는 근거가 된다.",
            "기술적 완성도 측면에서는 관계형 데이터베이스 기반 상태 관리, 트랜잭션 처리, 외부 검증 연동, 중복 지급 방지, 티켓 사용 이력 관리, 운영 권한 분리, 안전 관리 흐름이 결합되어 있다. 이는 단순 기능 구현을 넘어 실제 서비스 운영을 고려한 백엔드 설계라는 의미를 가진다.",
            "결론적으로 AirConnect의 작성 범위 백엔드 구현은 사용자 신뢰 확보, 개인 추천 품질 향상, 외부 검증 기반 결제 처리, 보상 안정성, 운영 제어, 안전 관리라는 핵심 요구사항을 충족한다. 향후에는 추천 조건의 세분화, 궁합 분석 품질 개선, 결제 상품 확장, 운영 정책 자동화, 안전 기능 고도화로 발전할 수 있다.",
        ],
        None,
    ),
]


def add_context_summary(doc):
    add_heading(doc, "", "작성 범위 기능 요약", 2)
    rows = [
        ["인증", "소셜 인증, 토큰 발급과 갱신, 탈퇴 흐름"],
        ["회원", "프로필 등록과 수정, 관심사, 취향 데이터, 차단 목록 관리"],
        ["개인 매칭", "관심사 기반 후보 추출, 요청 상태 관리, 중복 연결 방지"],
        ["궁합", "외부 인공지능 모델을 활용한 다항목 점수화와 요약 문장 생성"],
        ["결제", "서버 검증과 계정 식별값 기반 중복 지급 방지"],
        ["광고 보상", "전자 서명 검증과 보상 세션 기반 멱등 처리"],
        ["티켓", "지급, 차감, 잔액 일관성, 사용 이력 관리"],
        ["운영 관리", "운영 대시보드, 사용자와 신고 조회, 권한 분리"],
        ["공지", "노출 기간과 우선순위 기반 공지 관리"],
        ["점검", "점검 모드와 클라이언트 버전 제어"],
        ["안전", "신고 접수, 처리, 차단, 제재 정책"],
    ]
    add_table(doc, ["영역", "보고서에서의 해석"], rows, [1900, 7460])
    add_paragraph(
        doc,
        "위 기능들은 작성 범위에 포함되는 백엔드 기능이다. 본 보고서는 해당 기능들이 사용자 신뢰, 추천 품질, 비용 처리, 운영 제어, 서비스 안전성을 어떻게 구성하는지 분석한다.",
    )


def build():
    doc = Document()
    configure_document(doc)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "AirConnect 백엔드 기능 설계 및 구현 보고서", size=9, color="555555")

    add_cover(doc)
    add_toc(doc)

    for number, title, paragraphs, table in SECTIONS:
        add_heading(doc, number, title, 1)
        if number == "1":
            add_context_summary(doc)
        if number == "10":
            add_flow_diagrams(doc)
        for index, text in enumerate(paragraphs):
            add_paragraph(doc, text)
            if number == "3" and index == 0:
                add_project_erd(doc)
        if table:
            add_table(doc, table["headers"], table["rows"], table["widths"])
        if number in set():
            page_break(doc)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
