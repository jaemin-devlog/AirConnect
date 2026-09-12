from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE


OUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUT_DIR / "AirConnect_최종보고서_10페이지.docx"

FONT = "Nanum Gothic"
BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(31, 31, 31)
MUTED = RGBColor(95, 104, 117)
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F3F5F7"
BORDER = "D8DEE6"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.15):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_text(paragraph, text, size=10.5, bold=False, color=DARK):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return run


def add_para(doc, text, size=10.5, after=6, before=0, bold=False, color=DARK, align=None):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before, after=after)
    if align is not None:
        p.alignment = align
    add_text(p, text, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    if level == 1:
        size, before, after, color = 16, 0, 8, BLUE
    elif level == 2:
        size, before, after, color = 12.5, 8, 5, BLUE
    else:
        size, before, after, color = 11.5, 6, 4, RGBColor(45, 68, 92)
    set_paragraph_spacing(p, before=before, after=after, line=1.1)
    add_text(p, text, size=size, bold=True, color=color)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, before=0, after=3, line=1.12)
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.2)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    add_text(p, text, size=10, color=DARK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_paragraph_spacing(p, before=0, after=3, line=1.12)
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    add_text(p, text, size=10, color=DARK)
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(int(Cm(width_cm).twips)))
    cell.width = Cm(width_cm)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    hdr = table.rows[0].cells
    for idx, text in enumerate(headers):
        cell = hdr[idx]
        shade_cell(cell, header_fill)
        set_cell_margins(cell)
        set_cell_width(cell, widths[idx])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0, line=1.05)
        add_text(p, text, size=9.5, bold=True, color=BLUE)

    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cell = cells[idx]
            set_cell_margins(cell)
            set_cell_width(cell, widths[idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, after=0, line=1.08)
            if idx == 0 and len(headers) <= 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_text(p, str(text), size=9.2, bold=True, color=DARK)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_text(p, str(text), size=9.2, color=DARK)
    add_para(doc, "", after=2)
    return table


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="C9D8EA", size="6")
    cell = table.cell(0, 0)
    set_cell_width(cell, 16.8)
    set_cell_margins(cell, top=130, bottom=130, start=170, end=170)
    shade_cell(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=3, line=1.12)
    add_text(p, title, size=10.5, bold=True, color=BLUE)
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, after=0, line=1.12)
    add_text(p2, body, size=9.7, color=DARK)
    add_para(doc, "", after=2)


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text_run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    text_run.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text_run)
    run._r.append(end)


def configure_document(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.12

    heading_tokens = {
        "Heading 1": (16, BLUE, 0, 8),
        "Heading 2": (12.5, BLUE, 8, 5),
        "Heading 3": (11.5, RGBColor(45, 68, 92), 6, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1

    for style_name, size, color, bold in (
        ("AC Cover Title", 28, BLUE, True),
        ("AC Cover Subtitle", 18, DARK, True),
        ("AC Cover Meta", 12, MUTED, False),
    ):
        if style_name not in styles:
            style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = color
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.1

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(header, after=0, line=1)
    add_text(header, "AirConnect 최종 보고서", size=8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(footer, after=0, line=1)
    add_text(footer, "쪽 ", size=8.5, color=MUTED)
    add_page_number(footer)


def build_report():
    doc = Document()
    configure_document(doc)

    # 1
    cover = doc.add_paragraph(style="AC Cover Title")
    add_text(cover, "AirConnect", size=28, bold=True, color=BLUE)
    cover.paragraph_format.space_after = Pt(2)
    cover = doc.add_paragraph(style="AC Cover Subtitle")
    add_text(cover, "최종 프로젝트 보고서", size=18, bold=True, color=DARK)
    cover.paragraph_format.space_after = Pt(8)
    cover = doc.add_paragraph(style="AC Cover Meta")
    add_text(cover, "항공 콘셉트 기반 1:1 소개팅 및 다대다 과팅 매칭 서비스", size=12, color=MUTED)
    cover.paragraph_format.space_after = Pt(28)
    add_callout(
        doc,
        "보고서 범위",
        "본 보고서는 AirConnect 백엔드 프로젝트의 전체 구현 범위를 정리한 최종 산출물이다. 코드 기준으로 확인된 인증, 사용자, 매칭, 그룹 매칭, 채팅, 알림, 결제, 운영자, 통계, 보안, 배포 기능을 중심으로 작성하였다."
    )
    add_heading(doc, "1. 프로젝트 개요", 1)
    add_para(doc, "AirConnect는 대학교 구성원을 대상으로 한 매칭 서비스이다. 서비스의 핵심 콘셉트는 ‘인연의 활주로’이며, 사용자는 항공권을 고르듯 가볍게 상대를 탐색하고, 마일리지에 해당하는 티켓을 사용해 관심 표현과 매칭 요청을 진행한다. 단순 소개팅 기능에 머무르지 않고 1:1 매칭, 팀 기반 과팅, 실시간 채팅, 알림, 신고·차단, 결제와 광고 보상, 관리자 운영 도구까지 갖춘 운영 가능한 백엔드 시스템을 목표로 설계되었다.")
    add_para(doc, "구현 관점에서 가장 큰 특징은 사용자 경험과 운영 안정성을 동시에 고려했다는 점이다. 사용자는 소셜 로그인 후 프로필을 완성하고, 추천 후보를 확인하고, 티켓을 사용해 매칭을 요청한다. 매칭이 수락되면 채팅방이 자동 생성되고, 그룹 매칭에서는 팀방 구성, 준비 상태 확인, 큐 진입, 최종 그룹 채팅방 생성으로 이어지는 별도 흐름이 제공된다. 운영자는 관리자 API를 통해 사용자 상태, 신고, 티켓, 공지, 채팅방, 퍼널 지표, 푸시 발송 상태를 확인할 수 있다.")
    add_heading(doc, "핵심 목표", 2)
    add_bullet(doc, "낮은 진입 장벽: 소셜 로그인과 단계적 온보딩으로 먼저 서비스를 둘러보고 필요한 시점에 정보를 보완한다.")
    add_bullet(doc, "신뢰 가능한 캠퍼스 매칭: 학교 이메일 인증, 신고·차단, 사용자 상태 제한, 프로필 검증을 통해 매칭 품질을 높인다.")
    add_bullet(doc, "실시간 소통: STOMP WebSocket과 Redis Pub/Sub 기반으로 채팅과 그룹 매칭 이벤트를 빠르게 전달한다.")
    add_bullet(doc, "운영 가능성: 관리자 대시보드, 감사 로그, 알림 outbox, API 사용 통계, 유지보수 모드로 실제 서비스 운영을 고려한다.")

    add_page_break(doc)

    # 2
    add_heading(doc, "2. 전체 아키텍처와 기술 스택", 1)
    add_para(doc, "AirConnect는 Spring Boot 기반의 REST API 서버이며, 실시간 통신에는 WebSocket STOMP를 함께 사용한다. 데이터 저장은 JPA 엔티티와 Repository 중심으로 구성되어 있고, Redis는 인증·채팅 세션·학교 이메일 인증 코드·그룹 매칭 큐와 같은 짧은 수명 상태 관리에 사용된다. 외부 연동은 Kakao/Apple 소셜 인증, Apple/Google 인앱 결제 검증, Firebase Cloud Messaging, AdMob 보상형 광고 SSV, 선택적 OpenAI 궁합 요약으로 구성된다.")
    add_table(
        doc,
        ["영역", "사용 기술", "적용 내용"],
        [
            ("언어/런타임", "Java 17", "Spring Boot 3.4.3 기반 애플리케이션 구현"),
            ("웹/API", "Spring Web, Validation", "REST 컨트롤러, DTO 검증, 공통 응답 및 예외 처리"),
            ("데이터", "Spring Data JPA, MySQL", "사용자, 매칭, 채팅, 결제, 알림, 운영 데이터 영속화"),
            ("캐시/상태", "Redis", "STOMP 세션, 인증 코드, 큐 토큰, 활동 기록 등 단기 상태 관리"),
            ("보안", "Spring Security, JWT, BCrypt", "Stateless 인증, 관리자 권한, 토큰 갱신, 비밀번호 해싱"),
            ("실시간", "WebSocket STOMP, Redis Pub/Sub", "채팅 메시지, 구독 권한, 읽음 상태, 그룹 매칭 이벤트"),
            ("알림", "Firebase Admin SDK", "FCM 푸시 발송, 디바이스 토큰 관리, outbox 재시도"),
            ("수익화", "Apple/Google IAP, AdMob SSV", "티켓 구매 검증, 환불 처리, 광고 보상 지급"),
            ("배포", "Docker, Docker Compose, GitHub Actions", "JDK 17 이미지 빌드, MySQL/Redis 포함 배포 구성"),
        ],
        [3.0, 4.5, 9.3],
    )
    add_heading(doc, "모듈 구조", 2)
    add_para(doc, "패키지는 도메인별로 분리되어 있다. 대표 경계는 `auth`, `user`, `matching`, `groupmatching`, `chat`, `notification`, `iap`, `ads`, `moderation`, `admin`, `analytics`, `statistics`, `maintenance`, `global`이며, 각 모듈은 Controller, Service, Repository, DTO, Entity, ErrorCode를 나누어 책임을 분리한다. 대표 API는 인증 `/api/v1/auth/**`, 사용자 `/api/v1/users/**`, 1:1 매칭 `/api/v1/matching/**`, 그룹 매칭 `/api/v1/matching/team-rooms/**`, 채팅 `/api/v1/chat/**`와 `/ws-stomp`, 운영자 `/api/v1/admin/**`로 구성된다.", after=2)

    add_page_break(doc)

    # 3
    add_heading(doc, "3. 인증, 온보딩, 사용자 프로필", 1)
    add_para(doc, "인증 모듈은 소셜 로그인과 관리자 로그인을 분리하여 제공한다. 일반 사용자는 Kakao 또는 Apple 인증 토큰을 제출하고, 서버는 각 제공자의 클라이언트를 통해 socialId를 조회한 뒤 사용자 계정을 생성하거나 기존 계정을 찾는다. Apple 로그인은 이메일을 별도로 해석하며, 신규 계정 생성 시 디바이스 바인딩을 확인해 동일 기기에서 여러 소셜 계정을 남발하는 위험을 줄인다.")
    add_para(doc, "JWT는 access token과 refresh token으로 분리된다. refresh token에는 deviceId가 포함되고, 저장 시에는 TokenHashService를 통해 해시값으로 보관한다. 갱신 요청에서는 토큰의 기기 정보와 요청 deviceId를 비교하고, 저장 토큰과 맞지 않으면 재사용 탐지로 간주해 해당 refresh token을 삭제한다. 관리자 로그인은 이메일·비밀번호 기반이며 BCrypt 검증, 로그인 시도 제한, 전체 refresh token 폐기 흐름을 포함한다.")
    add_heading(doc, "사용자 온보딩", 2)
    add_para(doc, "회원가입은 이미 소셜 로그인으로 만들어진 사용자에게 이름, 닉네임, 학번, 학과, 키, 나이, MBTI, 흡연 여부, 성별, 군필 여부, 종교, 거주지, 자기소개, 인스타그램 정보를 채우는 과정이다. 완료 시 사용자 상태와 온보딩 상태가 갱신되고, AnalyticsService에 `SIGN_UP_COMPLETED` 서버 이벤트가 저장된다.")
    add_bullet(doc, "내 정보 조회는 사용자 기본 정보, 프로필 존재 여부, 프로필 이미지 업로드 여부, 학교 이메일 인증 여부, 티켓 잔액, iOS appAccountToken까지 반환한다.")
    add_bullet(doc, "닉네임 변경은 최대 길이와 정규화 처리를 수행하고, 기존 닉네임과 다를 경우 14일 쿨다운을 적용한다.")
    add_bullet(doc, "프로필 이미지는 확장자, MIME, 실제 이미지 디코딩, 픽셀 수, 파일 크기를 검증하고, 메타데이터 제거 재인코딩 후 안전한 UUID 파일명으로 저장한다.")
    add_bullet(doc, "학교 이메일 인증은 `office.hanseo.ac.kr` 도메인만 허용하고, Redis에 5분짜리 인증 코드를 저장한다. 재전송 쿨다운, IP/이메일 기준 시도 제한, 인증 완료 토큰 발급을 포함한다.")
    add_heading(doc, "마일스톤 보상", 2)
    add_para(doc, "프로필 이미지 업로드와 학교 이메일 인증은 선택 보상형 마일스톤으로 설계되어 있다. 사용자가 기능을 완료하면 UserMilestone 엔티티가 생성되고 보상 티켓이 지급된다. 이 구조는 필수 차단보다 긍정적 보상으로 프로필 신뢰도를 높이는 방식이며, VerifiedSchoolEmailFilter도 실제 접근 차단 대신 상태와 보상 로직에 집중하도록 구현되어 있다.")
    add_callout(doc, "구현 포인트", "사용자 기능은 단순 CRUD가 아니라 인증 상태, 사용자 상태, 이미지 보안, 중복 보상 방지, Apple 계정 해지 연동, 채팅 세션 무효화까지 고려한다. 특히 탈퇴·정지·삭제 상태를 여러 서비스에서 공통으로 검사해 비정상 사용자가 매칭과 채팅에 접근하지 못하게 한다.")

    add_page_break(doc)

    # 4
    add_heading(doc, "4. 1:1 매칭과 궁합 분석", 1)
    add_para(doc, "1:1 매칭은 AirConnect의 핵심 기능이다. 사용자는 추천 후보를 새로고침하고, 마음에 드는 상대에게 매칭 요청을 보낸다. 추천은 기본적으로 성별 조건을 고려하며, 별도 API로 동일 성별 추천도 지원한다. 추천 후보는 한 번 노출된 사용자를 MatchingExposure에 기록하여 새로고침 때 같은 후보가 반복되는 문제를 줄이고, 모든 후보가 소진되면 노출 이력을 초기화하여 다시 순환한다.")
    add_para(doc, "티켓 경제도 매칭 흐름과 연결되어 있다. 추천 후보가 2명 이상 반환될 때 추천 새로고침 비용으로 1티켓을 차감하고, 실제 매칭 요청은 2티켓을 차감한다. 차감 시점은 검증이 모두 끝난 뒤로 미루어져 있으며, TicketLedger에 소비 내역을 남겨 운영자가 나중에 잔액 변동을 추적할 수 있다. 후보가 부족하여 1명 이하만 반환되면 티켓을 차감하지 않는 예외도 구현되어 사용자 불만을 줄인다.")
    add_heading(doc, "매칭 요청 흐름", 2)
    add_number(doc, "요청자와 대상 사용자가 같은지 확인하고, 양쪽 사용자의 ACTIVE 상태와 프로필 성별 정보를 검증한다.")
    add_number(doc, "신고·차단 정책을 조회해 서로 차단 관계이면 매칭 요청과 채팅 생성을 막는다.")
    add_number(doc, "해당 대상이 실제 추천 응답에 포함되어 노출된 후보인지 MatchingExposure로 검증한다.")
    add_number(doc, "이미 ACCEPTED 연결이 있고 채팅방이 살아 있으면 기존 채팅방을 반환하고, PENDING이면 중복 요청을 거부한다.")
    add_number(doc, "REJECTED 또는 종료된 ACCEPTED는 재요청 가능하도록 PENDING으로 전환하고, 신규 요청은 MatchingConnection을 생성한다.")
    add_number(doc, "상대방에게 매칭 요청 알림을 생성하고, `MATCH_REQUEST_SENT` 분석 이벤트를 저장한다.")
    add_heading(doc, "수락과 채팅방 생성", 2)
    add_para(doc, "상대가 요청을 수락하면 MatchingConnection 상태가 ACCEPTED로 전환되고, ChatService를 통해 1:1 PERSONAL 채팅방이 생성된다. 같은 connectionId 또는 같은 사용자 쌍으로 이미 채팅방이 존재하면 중복 생성하지 않고 기존 방을 재사용한다. 이 덕분에 네트워크 재시도나 동시 요청 상황에서도 사용자가 같은 상대와 여러 채팅방을 갖는 문제가 줄어든다.")
    add_heading(doc, "궁합 분석", 2)
    add_para(doc, "Compatibility 모듈은 두 사용자의 프로필을 기반으로 궁합 점수와 설명을 제공한다. 나이, 학과, 학번, 키, MBTI, 흡연 여부, 종교, 거주지를 요소별로 점수화하고, 원점수 25~100 범위를 사용자에게 보기 좋은 50~100 점수로 변환한다. MBTI는 별도 CompatibilityTable을 통해 이상적 조합, 가능성 있는 조합, 신중한 조합으로 구분한다. OpenAI 연동이 켜져 있으면 결과 요약을 생성하고, 비활성화되거나 실패하면 로컬 fallback 문구를 사용한다.")

    add_page_break(doc)

    # 5
    add_heading(doc, "5. 다대다 과팅과 그룹 매칭", 1)
    add_para(doc, "그룹 매칭은 팀을 먼저 만들고 상대 팀과 매칭되는 구조이다. 방장은 팀명, 팀 성별, 팀 규모, 상대 성별 필터, 공개 여부를 설정해 임시 팀방을 만든다. 팀방 생성과 동시에 그룹 채팅방이 생성되고, 방장 멤버십과 준비 상태가 같은 트랜잭션 안에서 저장된다. 공개방은 모집 목록에 노출되고, 비공개방은 초대 코드로 입장할 수 있다.")
    add_para(doc, "임시 팀방은 OPEN, READY_CHECK, QUEUE_WAITING, MATCHED 같은 상태를 가진다. 구성원은 공개방 참여, 초대 코드 참여, 퇴장, 방장에 의한 추방이 가능하며, 멤버 변화가 생기면 준비 상태를 초기화하여 오래된 ready 값으로 큐에 들어가는 문제를 막는다. 팀 성별 검증, 동일 사용자의 다중 활성 팀방 방지, 팀명 중복 방지, 방장 전용 액션 검증도 포함되어 있다.")
    add_heading(doc, "준비와 큐 처리", 2)
    add_para(doc, "팀원이 모두 모이면 ready 상태를 통해 매칭 큐에 진입한다. 큐 처리는 Redis 기반 토큰과 짧은 TTL의 프로세스 락을 사용해 동시에 여러 워커가 같은 팀을 매칭하는 문제를 줄인다. `MATCH_FINALIZATION_DELAY`로 최종화 지연을 두어 이벤트 전달과 상태 정리가 급하게 엉키지 않도록 설계되어 있다. 매칭이 성사되면 두 팀의 구성원을 합쳐 최종 그룹 채팅방을 생성하고, 임시방에는 최종 방 이동 안내 시스템 메시지를 발행한다.")
    add_table(
        doc,
        ["구분", "구현 내용", "효과"],
        [
            ("팀방 생성", "팀명, 팀 규모, 팀 성별, 상대 필터, 공개 여부 입력", "사용자가 과팅 팀을 자율적으로 구성"),
            ("입장 방식", "공개 모집 목록 또는 초대 코드 참여", "공개 모집과 지인 초대 모두 지원"),
            ("멤버 관리", "퇴장, 추방, 준비 상태 초기화", "방장 권한과 팀 상태 일관성 확보"),
            ("큐 진입", "전체 ready 후 Redis 큐 및 락 기반 처리", "중복 매칭과 동시성 위험 감소"),
            ("최종방", "매칭 성공 후 GROUP 채팅방 자동 생성", "과팅 성사 직후 소통 연결"),
        ],
        [2.6, 7.0, 7.2],
    )
    add_heading(doc, "실시간 이벤트", 2)
    add_para(doc, "GMatchingEventPublisher와 GMatchingPushService는 팀방 상태, 멤버 변화, 준비 상태, 매칭 결과를 실시간으로 전달하는 역할을 한다. STOMP 구독 권한은 해당 팀방 접근 가능 여부를 검사한 뒤 허용되므로, 팀원이 아닌 사용자가 임의로 팀방 이벤트를 구독하는 위험을 줄인다. 이 구조는 프론트엔드가 팀방 화면에서 구성원 수, 준비 상태, 매칭 진행 상태를 즉시 갱신할 수 있게 한다.")

    add_page_break(doc)

    # 6
    add_heading(doc, "6. 실시간 채팅 시스템", 1)
    add_para(doc, "채팅 모듈은 1:1 PERSONAL 채팅방과 GROUP 채팅방을 모두 지원한다. 1:1 채팅방은 매칭 수락과 연결되어 자동 생성되며, 그룹 채팅방은 임시 팀방과 최종 과팅방에서 사용된다. ChatService는 채팅방 생성, 멤버 추가, 입장·퇴장 시스템 메시지, 메시지 저장과 발행, 방 목록 조회, 참여자 프로필 조회, 읽음 처리, 메시지 삭제를 담당한다.")
    add_para(doc, "실시간 전송은 WebSocket STOMP로 처리된다. 클라이언트는 `/ws-stomp`에 연결하고 `/pub/chat/message` 형태로 메시지를 발행하며, 방 구독은 `/sub/chat/room/{roomId}`로 이루어진다. StompHandler는 CONNECT 시 Authorization Bearer 토큰을 검증해 Principal을 설정하고, SUBSCRIBE 시 채팅방 멤버인지 확인한다. 구독 권한이 없는 사용자는 AccessDeniedException으로 차단된다.")
    add_heading(doc, "Redis와 세션 관리", 2)
    add_para(doc, "채팅은 Redis Pub/Sub과 세션 저장을 함께 활용한다. 사용자가 STOMP에 접속하면 sessionId와 userId를 Redis에 저장하고, 방 구독 시 sessionId, subscriptionId, roomId의 매핑을 기록한다. 사용자가 방을 보고 있으면 읽음 상태를 동기화하고, 메시지 발송 시 해당 방의 세션 정보를 활용해 방 목록 업데이트와 읽음 상태 반영을 수행할 수 있다. 탈퇴 시에는 특정 사용자의 활성 STOMP 세션을 찾아 무효화하는 기능도 포함되어 있다.")
    add_heading(doc, "사용자 보호와 편의 기능", 2)
    add_bullet(doc, "차단 관계인 사용자는 1:1 채팅방 생성과 메시지 흐름에서 제한된다.")
    add_bullet(doc, "채팅방 목록은 마지막 메시지, 안 읽은 메시지 수, 참여자 정보와 함께 제공되어 모바일 앱 UI에 바로 쓰기 좋다.")
    add_bullet(doc, "메시지 삭제는 물리 삭제보다 삭제 플래그를 활용하는 방식으로 구현되어 운영 추적 가능성을 남긴다.")
    add_bullet(doc, "1:1 상대 프로필, 그룹 참여자 프로필, 특정 참여자 프로필 조회 API가 분리되어 화면별 필요한 정보를 가져올 수 있다.")
    add_callout(doc, "설계 의의", "채팅은 단순 메시지 저장소가 아니라 매칭 결과의 후속 경험이다. 따라서 매칭 연결, 그룹 팀방, 차단 정책, 알림, 읽음 상태, STOMP 보안을 하나로 연결해 실제 사용자 흐름에서 끊김이 없도록 구성되어 있다.")

    add_page_break(doc)

    # 7
    add_heading(doc, "7. 알림, 공지, 푸시 전달", 1)
    add_para(doc, "알림 모듈은 앱 내 알림함과 외부 푸시 발송을 분리한다. NotificationService는 먼저 알림 원본을 저장하고, 사용자 알림 설정과 방해금지 시간을 확인한 뒤 푸시가 허용되면 NotificationOutbox 행을 생성한다. 이 방식은 비즈니스 트랜잭션과 외부 FCM 호출을 분리하는 outbox 패턴에 가깝다. 알림 원본 저장이 성공하면 사용자는 앱 내 알림함에서 확인할 수 있고, 푸시 발송 실패는 별도 outbox 상태로 추적된다.")
    add_heading(doc, "알림 설정", 2)
    add_para(doc, "NotificationPreferenceService는 사용자별 기본 설정을 자동 생성하고, 푸시, 인앱 알림, 매칭 요청, 매칭 결과, 그룹 매칭, 채팅 메시지, 마일스톤, 약속 리마인더 항목을 분리해 관리한다. 방해금지 시간은 사용자의 timezone을 기준으로 계산되며, 앱 내 알림 노출은 유지하되 푸시 발송만 제한할 수 있다.")
    add_heading(doc, "푸시 디바이스와 outbox", 2)
    add_para(doc, "PushDeviceService는 userId와 deviceId 기준으로 디바이스를 upsert하고, FCM pushToken 소유권이 다른 디바이스로 이동하면 기존 소유권을 해제한다. NotificationOutboxWorker는 주기적으로 PENDING outbox를 claim하고 발송한다. 성공 시 SENT, 무효 토큰은 SKIPPED와 토큰 비활성화, 일시 오류는 1분·5분·30분 간격 재시도, 최종 실패는 FAILED 상태로 전이한다. 오래 PROCESSING에 묶인 행은 복구 작업으로 다시 처리 가능하게 만든다.")
    add_table(
        doc,
        ["알림 유형", "대표 발생 시점", "전달 방식"],
        [
            ("MATCH_REQUEST_RECEIVED", "1:1 매칭 요청 수신", "인앱 알림 + FCM 푸시"),
            ("MATCH_REQUEST_ACCEPTED/REJECTED", "상대의 수락 또는 거절", "인앱 알림 + 매칭 결과 설정 반영"),
            ("GROUP_MATCHED", "그룹 매칭 성공", "팀방 이벤트 + 푸시"),
            ("CHAT_MESSAGE_RECEIVED", "채팅 메시지 수신", "방 목록 업데이트 + 푸시"),
            ("MILESTONE_REWARDED", "프로필 이미지/학교 이메일 보상", "보상 안내 알림"),
            ("SYSTEM_ANNOUNCEMENT", "관리자 공지 발송", "공지/알림함/푸시"),
        ],
        [4.2, 6.3, 6.3],
        header_fill=LIGHT_GRAY,
    )
    add_heading(doc, "공지 기능", 2)
    add_para(doc, "NoticeService는 관리자가 발행한 공지를 최신순으로 제공한다. 관리자 공지 발송은 AdminNotice 엔티티에 기록되고, 대상 사용자 수와 활성 사용자 필터를 반영한다. 공지는 운영팀이 서비스 장애, 점검, 정책 변경, 이벤트성 보상을 사용자에게 전달하는 기본 채널로 활용된다.")

    add_page_break(doc)

    # 8
    add_heading(doc, "8. 티켓 경제, 인앱 결제, 광고 보상", 1)
    add_para(doc, "AirConnect의 티켓은 추천 새로고침, 매칭 요청, 보상형 광고, 인앱 결제와 연결된 핵심 경제 단위이다. 사용자 엔티티는 티켓 잔액을 갖고, 모든 지급·소비는 TicketLedger를 통해 추적된다. 이 구조는 사용자의 잔액 변동을 관리자 화면에서 확인하고, 중복 지급·중복 소비와 같은 문제를 감사할 수 있게 한다.")
    add_heading(doc, "인앱 결제", 2)
    add_para(doc, "IAP 모듈은 Apple과 Google Play 결제를 모두 고려한다. iOS는 signedTransactionInfo와 transactionId, appAccountToken을 검증하고, Android는 productId, purchaseToken, orderId, packageName을 검증한다. StoreVerifierResolver가 스토어별 검증기를 선택하고, IapProcessingService가 검증 결과를 주문 처리 흐름으로 넘긴다.")
    add_para(doc, "주문 처리는 멱등성을 중요하게 다룬다. Apple은 transactionId, Google은 purchaseToken 기준으로 기존 주문을 찾고, 없으면 PENDING 주문을 생성한 뒤 `findByIdForUpdate`로 잠근다. 이미 GRANTED인 주문은 중복 지급하지 않고 기존 결과를 반환한다. 취소·환불·revoked 상태가 확인되면 REVOKED 또는 REFUNDED로 전환하고, 이미 지급된 티켓은 IapRefundService를 통해 환불 처리한다.")
    add_table(
        doc,
        ["상품", "productId", "지급 티켓"],
        [
            ("Economy", "AirConnect_Economy_5", "5"),
            ("Premium Economy", "AirConnect_PremiumEconomy_10", "12"),
            ("Business", "AirConnect_Business_30", "30"),
            ("First Class", "AirConnect_FirstClass_50", "70"),
            ("Legacy Packs", "com.airconnect.tickets.pack*", "5~70"),
        ],
        [3.5, 8.7, 4.6],
    )
    add_heading(doc, "광고 보상", 2)
    add_para(doc, "AdReward 모듈은 보상형 광고 시청 후 AdMob SSV(Server Side Verification) 콜백을 처리한다. 사용자는 먼저 보상 세션을 생성하고, AdMob은 callback URL에 transaction_id, custom_data, signature, key_id를 포함해 호출한다. 서버는 custom_data에서 sessionKey를 추출하고, Tink 기반 AdmobSignatureVerifier로 서명을 검증한 뒤 세션 상태와 만료 여부, transactionId 중복 여부를 확인한다.")
    add_para(doc, "광고 보상 지급도 멱등적으로 구현되어 있다. 이미 지급된 세션이면 기존 TicketLedger를 찾아 ALREADY_GRANTED 응답을 반환하고, 신규 지급이면 사용자 행을 잠근 뒤 티켓을 더하고 AD_REWARD_SESSION refType의 ledger를 남긴다. 이 구조는 AdMob 콜백 재시도, 네트워크 중복, 운영자 확인 상황에서 티켓 경제의 일관성을 유지한다.")

    add_page_break(doc)

    # 9
    add_heading(doc, "9. 운영자, 통계, 신고·차단, 점검 모드", 1)
    add_para(doc, "운영 기능은 AirConnect가 실제 서비스로 동작하기 위해 중요한 부분이다. AdminController는 사용자 목록, 상세 조회, 상태 조치, 영구 삭제, 매칭 기록, 신고 목록, 신고 상태 변경, 티켓 잔액과 ledger, 수동 티켓 조정, 통계, 공지, 채팅방 조회, 운영 대시보드, 감사 로그를 제공한다. Spring Security 설정에서 `/api/v1/admin/**`은 ADMIN 역할만 접근할 수 있다.")
    add_heading(doc, "사용자와 신고 관리", 2)
    add_para(doc, "관리자는 사용자에게 SUSPEND, DELETE, REACTIVATE, RESTRICT_MATCHING, CLEAR_MATCHING_RESTRICTION 액션을 적용할 수 있다. 정지 또는 매칭 제한에는 사유와 해제 예정일을 넣을 수 있고, 사용자에게 운영 공지 알림이 전달된다. UserReportService는 신고 생성 시 신고자와 피신고자 존재 여부, 자기 신고 금지, 중복 신고 윈도우를 확인한다. UserBlockService와 UserBlockPolicyService는 차단 관계를 매칭과 채팅 정책에 반영한다.")
    add_heading(doc, "운영 대시보드", 2)
    add_para(doc, "AdminOperationsService는 운영 요약, API 사용 통계, 알림 outbox 모니터링, 1:1 매칭 퍼널, 그룹 매칭 퍼널, 알림 운영 통계, 운영 관리, 무결성 리포트를 제공한다. 예를 들어 1:1 매칭 퍼널은 추천 새로고침, 매칭 요청, 요청 수락, 거절/만료, 채팅방 생성 단계를 집계한다. 그룹 매칭 퍼널은 팀방 생성, 팀방 참여, 준비 완료 팀, 큐 진입, 매칭 성공, 최종 그룹 채팅방 생성으로 이어진다.")
    add_table(
        doc,
        ["운영 지표", "산출 근거", "활용"],
        [
            ("가입/활동", "사용자 수, 온보딩 완료, 최근 1/7/30일 활동", "성장성과 잔존율 확인"),
            ("매칭", "ACCEPTED 연결, 그룹 최종방, 퍼널 이벤트", "추천·요청·성사 병목 확인"),
            ("채팅", "삭제되지 않은 메시지 수, 채팅방 상세", "실사용 소통량 파악"),
            ("신고", "OPEN/IN_REVIEW 신고 수", "운영 리스크 우선순위 선정"),
            ("푸시", "outbox 상태, 실패, 평균 전달 시간", "알림 장애와 토큰 품질 확인"),
            ("API", "엔드포인트별 호출 수, 평균 응답 시간", "성능 최적화 대상 선정"),
        ],
        [3.5, 6.7, 6.6],
        header_fill=LIGHT_GRAY,
    )
    add_heading(doc, "점검 모드와 통계", 2)
    add_para(doc, "MaintenanceService는 단일 MaintenanceSetting을 사용해 전체 점검 상태, 제목, 메시지를 관리한다. MaintenanceModeFilter는 인증 필터 뒤에서 동작해 점검 중 일반 요청을 차단할 수 있고, `/api/v1/maintenance`는 공개되어 클라이언트가 점검 정보를 확인할 수 있다. StatisticsService는 메인 화면용 누적 가입자, DAU, 성비, 1:1·그룹 매칭 성사 합계, 요청 많은 학과 TOP 5를 제공한다.")

    add_page_break(doc)

    # 10
    add_heading(doc, "10. 품질, 보안, 검증, 향후 개선", 1)
    add_para(doc, "AirConnect는 기능 수가 많은 만큼 공통 품질 요소가 중요하다. GlobalExceptionHandler는 validation, 인증, 사용자, 매칭, 이메일 인증, IAP, 광고, 신고 등 도메인별 예외를 공통 ApiResponse와 ErrorBody로 변환한다. TraceIdFilter는 모든 요청에 X-Trace-Id를 부여해 로그와 응답을 연결한다. ApiRequestLoggingInterceptor는 `/api/v1/` 요청의 메서드, 경로, 상태 코드, 처리 시간을 저장해 운영 통계와 API 사용 분석의 기반을 만든다.")
    add_heading(doc, "보안 설계", 2)
    add_bullet(doc, "Spring Security는 stateless 세션 정책을 사용하고, JWT 필터를 통해 사용자 인증을 처리한다.")
    add_bullet(doc, "관리자 API는 ADMIN 역할만 접근할 수 있으며, Swagger 문서도 관리자 권한으로 제한된다.")
    add_bullet(doc, "Refresh token은 deviceId와 묶고 해시로 저장하여 토큰 탈취와 재사용 위험을 줄인다.")
    add_bullet(doc, "STOMP CONNECT와 SUBSCRIBE 단계에서 access token, 사용자 활성 상태, 채팅방/팀방 멤버십을 검증한다.")
    add_bullet(doc, "프로필 이미지는 실제 디코딩과 재인코딩으로 위장 파일, 과대 이미지, 메타데이터 노출 위험을 낮춘다.")
    add_heading(doc, "테스트와 검증 범위", 2)
    add_para(doc, "테스트 패키지에는 AuthServiceTest와 AuthServiceRaceTest, MatchingServiceTest와 MatchingServiceRaceTest, GMatchingServiceTest, GMatchingQueueWorkerTest, ChatServiceTest, IapProcessingServiceTest와 RaceTest, ModerationServiceTest와 RaceTest, NotificationInboxServiceTest, PushDeviceServiceTest, MaintenanceModeFilterTest, GlobalExceptionHandlerTest 등이 포함되어 있다. 이는 단순 성공 케이스뿐 아니라 동시성, 보안 필터, 예외 처리, 환불, outbox, 관리자 기능을 검증하려는 구조이다.")
    add_heading(doc, "향후 개선 제안", 2)
    add_bullet(doc, "CI에서 현재 build 단계가 `-x test`로 실행되므로, 배포 전 핵심 테스트 또는 smoke test를 포함하도록 개선한다.")
    add_bullet(doc, "관리자 통계 API의 주요 지표를 프론트엔드 대시보드와 연결하고, 알림 outbox 실패 알람을 운영 채널로 전송한다.")
    add_bullet(doc, "매칭 추천 로직에 프로필 완성도, 최근 활동성, 차단·신고 이력, 궁합 점수를 조합한 랭킹 모델을 추가한다.")
    add_bullet(doc, "학교 이메일 인증 완료 사용자에게 배지를 노출하고, 신고가 누적된 사용자의 추천 노출을 자동 제한하는 운영 정책을 보완한다.")
    add_bullet(doc, "인앱 결제와 광고 보상의 ledger를 관리자 화면에서 CSV로 내보낼 수 있게 하여 정산과 CS 대응을 강화한다.")
    add_para(doc, "종합하면 AirConnect는 캠퍼스 매칭이라는 서비스 목표를 중심으로, 매칭과 채팅의 사용자 경험, 티켓 기반 수익화, 신고·차단과 학교 인증을 통한 신뢰 장치, 관리자 운영과 통계까지 구현한 백엔드 프로젝트이다. 특히 단순 API 목록이 아니라 사용자 흐름, 운영 흐름, 외부 연동, 동시성 제어가 함께 설계되어 있어 실제 앱 서비스로 확장할 수 있는 기반을 갖추었다.")

    doc.core_properties.title = "AirConnect 최종 프로젝트 보고서"
    doc.core_properties.subject = "항공 콘셉트 기반 매칭 서비스 백엔드 구현 보고서"
    doc.core_properties.author = "AirConnect Team"
    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    path = build_report()
    print(path)
