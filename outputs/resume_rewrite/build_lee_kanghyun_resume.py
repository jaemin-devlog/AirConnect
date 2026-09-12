from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Cm, Inches, Pt, RGBColor


OUT_DIR = Path("/Users/rhee/AirConnect2/outputs/resume_rewrite")
PHOTO_PATH = OUT_DIR / "kanghyun_photo_crop.png"
DOCX_PATH = OUT_DIR / "이강현_백엔드_이력서_초안.docx"

BLUE = RGBColor(31, 67, 166)
DARK = RGBColor(17, 24, 39)
BODY = RGBColor(55, 65, 81)
FONT = "Malgun Gothic"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_no_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    tc_pr.append(borders)


def set_cell_margins(cell, top=70, start=95, bottom=70, end=95):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_fixed_width(table, widths_cm):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            cell = row.cells[idx]
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width / 2.54 * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F43A6")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)
    r_pr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    r_pr.append(size)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_section_heading(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text.upper())
    set_run_font(run, size=9.5, bold=True, color=BLUE)
    run.font.letter_spacing = 60 if hasattr(run.font, "letter_spacing") else None
    bottom = OxmlElement("w:pBdr")
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "8")
    border.set(qn("w:space"), "5")
    border.set(qn("w:color"), "E1E6EF")
    bottom.append(border)
    paragraph._p.get_or_add_pPr().append(bottom)
    return paragraph


def add_body_paragraph(doc, text, size=9.35, after=4.0, line=1.26):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=BODY)
    return paragraph


def add_label_value(paragraph, label, value):
    r1 = paragraph.add_run(label)
    set_run_font(r1, size=8.6, bold=True, color=DARK)
    r2 = paragraph.add_run(value)
    set_run_font(r2, size=8.6, color=BODY)


def add_bullet(doc, text, size=8.55):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.24)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    paragraph.paragraph_format.space_after = Pt(3.2)
    paragraph.paragraph_format.line_spacing = 1.22
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=BODY)
    return paragraph


def add_kv_line(doc, label, value, size=9.1, after=2.2):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.18
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=size, bold=True, color=DARK)
    value_run = paragraph.add_run(value)
    set_run_font(value_run, size=size, color=BODY)
    return paragraph


def add_kv_link_line(doc, label, text, url, size=8.85, after=1.8):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.18
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=size, bold=True, color=DARK)
    add_hyperlink(paragraph, text, url)
    return paragraph


def add_project(doc, title, meta_rows, overview, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11.1, bold=True, color=DARK)

    for label, value in meta_rows:
        add_kv_line(doc, f"{label} : ", value, size=9.0, after=2.2)
    add_kv_line(doc, "서비스 : ", overview, size=9.0, after=4.0)

    label_p = doc.add_paragraph()
    label_p.paragraph_format.space_before = Pt(1)
    label_p.paragraph_format.space_after = Pt(2)
    label_r = label_p.add_run("주요 구현")
    set_run_font(label_r, size=8.9, bold=True, color=BLUE)
    for bullet in bullets:
        add_bullet(doc, bullet, size=8.95)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(9.3)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.22

    for style_name, size, color in (
        ("Heading 1", 13.5, BLUE),
        ("Heading 2", 11, BLUE),
        ("Heading 3", 10, RGBColor(31, 77, 120)),
    ):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    bullet = styles["List Bullet"]
    bullet.font.name = FONT
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    bullet.font.size = Pt(8.95)
    bullet.paragraph_format.space_after = Pt(3)
    bullet.paragraph_format.line_spacing = 1.22


def build_resume():
    doc = Document()
    configure_document(doc)

    header = doc.add_table(rows=1, cols=2)
    set_table_fixed_width(header, [14.3, 3.7])
    for cell in header.rows[0].cells:
        set_cell_no_border(cell)
        set_cell_margins(cell, top=0, bottom=0, start=0, end=0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    left = header.rows[0].cells[0]
    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    name = p.add_run("이강현")
    set_run_font(name, size=24, bold=True, color=DARK)
    p2 = left.add_paragraph()
    p2.paragraph_format.space_after = Pt(6)
    sub = p2.add_run("Lee Kanghyun / Backend Engineer")
    set_run_font(sub, size=11.2, bold=True, color=BLUE)

    contact = left.add_paragraph()
    contact.paragraph_format.space_after = Pt(1)
    add_label_value(contact, "Phone  ", "010-9130-6469    ")
    add_label_value(contact, "GitHub  ", "")
    add_hyperlink(contact, "github.com/kanghyun-e", "https://github.com/kanghyun-e")

    right = header.rows[0].cells[1]
    if PHOTO_PATH.exists():
        rp = right.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp.add_run().add_picture(str(PHOTO_PATH), width=Cm(2.7))

    add_section_heading(doc, "Introduction")
    add_body_paragraph(
        doc,
        "결제/광고 보상처럼 돈과 신뢰가 걸린 기능은 중복 처리와 상태 전이를 설계 단계에서 먼저 막는 백엔드 개발자입니다. "
        "AirConnect를 앱스토어/플레이스토어에 출시/운영하며 인앱 결제 검증/광고 보상 검증/매칭 대기열 복구/신고 제재/관리자 운영 지표를 직접 설계하고 구현했습니다.",
        size=9.35,
        after=4,
    )
    add_body_paragraph(
        doc,
        "외부 API와 AI 응답도 정상 응답을 믿는 코드가 아니라 검증 가능한 파이프라인으로 다룹니다. "
        "MoneyWay에서는 OpenAI 응답을 후보 제한과 형식/일정 슬롯/예산 검증으로 통제했고 모의 테스트 30건에서 데이터베이스 밖 장소 생성/슬롯 누락/예산 초과를 0건으로 만들었습니다.",
        size=9.35,
        after=4,
    )

    add_section_heading(doc, "Personal Information")
    add_kv_line(doc, "NAME : ", "이강현 / Lee Kanghyun")
    add_kv_line(doc, "Education : ", "한서대학교 항공소프트웨어공학과 / 2021.03 ~ 2027.02 졸업예정 / GPA 3.81 / 4.5")

    add_section_heading(doc, "Project")
    add_project(
        doc,
        "[AirConnect] 대학생 소셜 매칭 / 티켓 경제 운영 백엔드 개발",
        [
            ("일정", "2025.11 ~ 현재"),
            ("기술 스택", "Java / Spring Boot / JPA / QueryDSL / MySQL 8 / Redis 7 / WebSocket / Docker Compose / Nginx / GitHub Actions / Apple IAP / Google IAP / AdMob"),
            ("참여 인력/역할", "백엔드 리드"),
            ("운영 환경", "Gabia 단일 서버 / 2vCore / 8GB RAM / SSD 50GB / Ubuntu 22.04"),
            ("운영 단계", "정식 홍보 전 기능 검증 단계 / 누적 가입자 18명 / 온보딩 완료 11명"),
        ],
        "대학생 1:1/그룹 매칭 / 실시간 채팅 / 티켓 결제 / 광고 보상 / 신고 제재 / 관리자 운영 기능을 포함한 iOS/Android 앱 백엔드입니다.",
        [
            "클라이언트 영수증만 믿고 티켓을 지급하면 위조나 재전송에 취약한 문제를 앱스토어 / 플레이스토어 서버 검증과 결제 고유값 중복 확인으로 막았습니다. 같은 결제가 동시에 2건 들어와도 1건만 지급되고 1건은 이미 처리된 결제로 종료되도록 검증했습니다.",
            "애플 결제는 사용자 연결값 / 상품 정보 / 결제 환경 / 원거래 정보를 함께 확인해 다른 사용자 결제 / 잘못된 환경의 결제 / 환불 또는 취소된 결제가 티켓 지급으로 이어지지 않도록 거절 경로를 분리했습니다.",
            "광고 보상은 AdMob 콜백 서명 / 세션값 / 보상 거래값을 검증한 뒤 같은 보상 요청이 반복되어도 보상이 1회만 반영되도록 설계해 중복 보상을 막았습니다.",
            "Redis 기반 그룹 매칭 대기열은 재시작이나 만료 상황에서 데이터베이스 상태와 어긋날 수 있어 대기 중인 팀 목록을 기준으로 큐 순서와 연결 정보를 다시 복구하도록 구성했습니다.",
            "신고 / 차단 / 제재 흐름은 같은 신고 2건이 동시에 들어와도 1건만 저장되고 1건은 중복으로 거절되도록 검증해 운영자가 신고 접수부터 검토 / 제재 / 매칭 제외까지 한 흐름으로 관리할 수 있게 했습니다.",
            "관리자 화면에 API 사용량 / 1:1 매칭 퍼널 / 그룹 매칭 퍼널 / 알림 발송 상태 / 7개 데이터 정합성 점검 / 감사 로그를 제공했습니다. 운영 로그 1,047건 이상 / 최근 7일 API 호출 117건 / 24개 엔드포인트 기준으로 운영 이슈를 기능 단위로 추적할 수 있게 했습니다.",
        ],
    )

    doc.add_page_break()
    add_project(
        doc,
        "[MoneyWay] 제주 AI 여행 일정 생성 서비스 백엔드 개발",
        [
            ("일정", "2025.05 ~ 2025.11"),
            ("기술 스택", "Java 21 / Spring Boot 3.4.5 / Spring Data JPA / MySQL 8 / Docker / OpenAI API / TourAPI"),
            ("참여 인력/역할", "백엔드 개발"),
            ("성과", "한국관광공사 2025 관광데이터 활용 공모전 우수상"),
        ],
        "제주 관광지 데이터와 AI를 활용해 사용자 조건에 맞는 여행 일정을 생성하는 서비스입니다.",
        [
            "TourAPI 기반 제주 관광지 수집과 엑셀 업로드 기반 식당/카페 등록 흐름을 구성하고 카테고리 / 검색 / 페이징을 포함한 REST API 48개를 구현했습니다.",
            "OpenAI 응답이 데이터베이스에 없는 장소를 만들거나 일정 슬롯을 누락하고 예산을 초과할 수 있는 문제를 입력 후보 제한 / 응답 형식 검증 / 슬롯 검증 / 예산 검증으로 제어했습니다.",
            "후보 장소가 부족해 일정 생성이 실패하는 지역을 위해 탐색 반경을 10km에서 시작해 후보 부족 시 1.5배씩 최대 50km까지 확장하도록 구현했습니다.",
            "모의 테스트 30건에서 데이터베이스 밖 장소 생성 / 슬롯 누락 / 예산 초과를 모두 0건으로 검증했습니다.",
        ],
    )

    add_section_heading(doc, "Education & Activities")
    for item in [
        "2026.03 ~ 현재 / 멋쟁이사자처럼 한서대학교 14기 부대표",
        "2025.09 ~ 2025.11 / 한국관광공사 2025 관광데이터 활용 공모전 우수상",
        "2025.07 ~ 2025.08 / 멋쟁이사자처럼 전국 연합 해커톤 2차 진출",
        "2025.03 ~ 2025.12 / 멋쟁이사자처럼 한서대학교 13기 일반 부원",
        "2022.04.22 / 한국사능력검정시험 / 58-107657",
    ]:
        add_bullet(doc, item, size=8.9)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_resume()
